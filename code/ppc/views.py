from django.contrib.auth.models import User
from django.contrib.auth.forms import UserCreationForm
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.template.loader import render_to_string
from django.db.models import Prefetch, F
from django.utils.text import slugify
from .permissions import exigir_acesso_ao_curso 
from weasyprint import HTML
from django.http import HttpResponseForbidden, Http404
from .permissions import cursos_do_usuario
from ppc.testes.analisar_pdf_uma_chamada import analisar_pdf_adaptativo, ErroOpenRouter
from .models import Curso, CineBrasilCurso, PPC, DinamicaEAD, ComponenteCurricular, Bibliografia, BibliografiaReferencia, Apendice, VinculoUsuarioCurso , RelacaoComponente, MembroNDE, ComponenteNaMatriz, Chamado, ResumoCargaHorariaPPC, GrupoEquivalenciaComponentes, ItemEquivalencia, ReferenciaBibliografica
from .forms import ( ObjetivosForm, EditarPermissoesForm, CursoForm,
                    ApresentacaoForm, ExposicaoMotivosForm, PrincipiosForm,
                    ExpectativasForm, TccForm, EstagioForm, AtividadesComplementaresForm,
                    PoliticasIntegradaForm, AvaliacaoEnsinoForm, AvalicaoProjetoCursoForm,
                    QualificacaoForm, RequisitosLegaisForm, ApendiceForm, DinamicaEADForm, 
                    EstruturaCurricularForm, ComponenteCurricularForm, BibliografiaForm, RelacaoComponenteForm,
                    ReferenciasForm, MembroNDEForm, ImportarPDFForm, ImportarPPCModeloAntigoForm, LimitesCargaHorariaFormSet,
                    ComponenteNaMatrizForm, ResumoCargaHorariaPPCForm, VinculoUsuarioCursoForm, ComponenteMatrizReferenciaForm, BibliografiaReferenciaForm )
from django.db.models import Q
from .importacao import extrair_dados_pdf
from .importacao_modelos_antigos import ErroImportacaoPPC, preparar_importacao_modelo_antigo
import logging
from functools import wraps
from types import SimpleNamespace
from django.db import transaction
import tempfile
from pathlib import Path
from io import BytesIO
try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:  # permite iniciar o sistema mesmo em ambientes sem o extra de exportação
    openpyxl = None
    Alignment = Border = Font = PatternFill = Side = None
logger = logging.getLogger(__name__)
from django.core.exceptions import ValidationError, PermissionDenied
import json
from django.http import JsonResponse
from django.views.decorators.http import require_POST, require_GET
from .models import MatrizReferenciaCurricular, ComponenteNaMatrizReferencia, EquivalenciaMatriz
from django.db.models import Count, Max
from .models import ImportacaoPendencia
from .services.resolucao_pendencias import resolver_pendencia_ambigua
from django.core.mail import send_mail
from config import settings
from django.utils import timezone
from .services.carga_horaria import sincronizar_carga_horaria_total
from django.db.models import Case, When, Value, IntegerField
from django.db.models import Q
from django.core.paginator import Paginator
from .models import ReferenciaBibliografica
from .services.fluxograma import montar_fluxograma_ppc, montar_fluxograma_referencia, copiar_relacoes_referencia_para_ppc, copiar_relacoes_entre_matrizes
from .forms import RelacaoComponenteReferenciaForm
from .models import MatrizReferenciaCurricular, RelacaoComponenteReferencia
from django.urls import reverse
from django.core.paginator import Paginator
from django.db.models import Count
from django.forms import formset_factory
from .forms import CodigoComponenteForm, CodigoComponenteFormSet
from django.utils.http import url_has_allowed_host_and_scheme
from .completude import confirmar_revisao
from ppc.services.eventos import registrar_evento, suspender_eventos_edicao
from ppc.services.regras_matriz import pendencias_referencias, relacoes_invalidas, contar_relacoes_invalidas, ppcs_fora_da_duracao_queryset, maior_periodo_ppc, validar_mudanca_periodo, periodo_relacao_invalido
from ppc.services.comparacao_bibliografias import calcular_operacoes_bibliografia, resumir_operacoes
from ppc.services.painel_bibliografias import (
    resumo_pendencias_bibliograficas, queryset_componentes_pendentes_ppc,
    queryset_componentes_pendentes_matriz, resumo_pendencias_matriz,
)
from ppc.services.carga_semanal import calcular_horas_semanais, formatar_horas_semanais_valor
from ppc.services.equivalencias_matrizes import (
    consultar_equivalencias_matriz, salvar_equivalencias_matriz,
    preparar_editor_equivalencias, candidatos_equivalencia_celula,
    previsualizar_equivalencias_matriz,
)
from ppc.services.equivalencias_ppc import preparar_editor_ppc, candidatos_ppc, salvar_ppc as salvar_equivalencias_ppc


def _bloquear_nde_em_detalhes_staff(view_func):
    """Mantém o fluxo staff_member_required e responde 403 a usuário logado NDE."""
    @wraps(view_func)
    def _wrapped(request, *args, **kwargs):
        if request.user.is_authenticated and not request.user.is_staff:
            raise PermissionDenied
        return view_func(request, *args, **kwargs)
    return _wrapped

@login_required
@require_POST
def confirmar_revisao_secao(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    confirmar_revisao(ppc, request.POST.getlist("secao"), request.user)
    messages.success(request, "Revisão confirmada.")
    destino = request.POST.get("next", "")
    if not url_has_allowed_host_and_scheme(
        destino, allowed_hosts={request.get_host()}, require_https=request.is_secure()
    ):
        destino = reverse("detalhe_curso", args=[ppc.curso_id])
    return redirect(destino)

@staff_member_required
def preencher_codigos_componentes(request):
    busca = request.GET.get('q', '').strip()
    base = ComponenteCurricular.objects.filter(codigo='')
    if busca:
        base = base.filter(nome__icontains=busca)
    base = base.annotate(
        n_ppcs=Count('uso_em_ppcs', distinct=True),
        n_matrizes=Count('uso_em_matrizes_referencia', distinct=True),
    ).order_by('nome', 'id')

    pagina = Paginator(base, 30).get_page(request.GET.get('page'))
    FormSet = formset_factory(CodigoComponenteForm, formset=CodigoComponenteFormSet, extra=0)

    if request.method == 'POST':
        formset = FormSet(request.POST)
        ids = [int(v) for k, v in request.POST.items()
               if k.endswith('-componente_id') and v.isdigit()]
        por_id = {c.id: c for c in base.filter(pk__in=ids)}
        if formset.is_valid():
            salvos = 0
            for form in formset:
                codigo = form.cleaned_data['codigo']
                componente = por_id.get(form.cleaned_data['componente_id'])
                if not codigo or componente is None:
                    continue
                componente.codigo = codigo
                componente.save(update_fields=['codigo'])
                salvos += 1
            messages.success(request, f"{salvos} código(s) salvo(s).")
            return redirect(request.get_full_path())
    else:
        componentes = list(pagina.object_list)
        por_id = {c.id: c for c in componentes}
        formset = FormSet(initial=[{'componente_id': c.id} for c in componentes])

    for form in formset:
        valor = form['componente_id'].value()
        form.obj = por_id.get(int(valor)) if valor not in (None, '') else None

    return render(request, 'ppc/preencher_codigos_componentes.html', {
        'formset': formset,
        'pagina': pagina,
        'busca': busca,
        'total_sem_codigo': pagina.paginator.count,
    })

@login_required
def fluxograma_ppc(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    return render(request, 'ppc/fluxograma.html', {
        'titulo': f"Fluxograma — {ppc.curso.nome}",
        'voltar_url': reverse('lista_componentes', args=[ppc.id]),
        'fluxograma': montar_fluxograma_ppc(ppc),
    })


@login_required
def fluxograma_matriz_referencia(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, id=matriz_id)
    return render(request, 'ppc/fluxograma.html', {
        'titulo': f"Fluxograma — {matriz}",
        'voltar_url': reverse('matriz_referencia_detalhe', args=[matriz.id]),
        'relacoes_url': reverse('relacoes_matriz_referencia', args=[matriz.id]),
        'fluxograma': montar_fluxograma_referencia(matriz),
    })


@login_required
def relacoes_matriz_referencia(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, id=matriz_id)
    form = RelacaoComponenteReferenciaForm(request.POST or None, matriz=matriz)
    if request.method == 'POST' and form.is_valid():
        try:
            form.save()
            messages.success(request, "Relação adicionada.")
            return redirect('relacoes_matriz_referencia', matriz_id=matriz.id)
        except ValidationError as exc:
            form.add_error(None, exc)
    relacoes = RelacaoComponenteReferencia.objects.filter(
        item__matriz=matriz
    ).select_related('item__componente', 'item_relacionado__componente')
    return render(request, 'ppc/relacoes_matriz_referencia.html', {
        'matriz': matriz, 'form': form, 'relacoes': relacoes,
    })


@login_required
@require_POST
def remover_relacao_referencia(request, relacao_referencia_id):
    relacao = get_object_or_404(RelacaoComponenteReferencia, id=relacao_referencia_id)
    matriz_id = relacao.item.matriz_id
    relacao.delete()
    messages.success(request, "Relação removida.")
    return redirect('relacoes_matriz_referencia', matriz_id=matriz_id)

@staff_member_required
def gerenciar_acessos(request):
    if request.method == 'POST':
        form = VinculoUsuarioCursoForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Acesso concedido.")
            return redirect('gerenciar_acessos')
    else:
        form = VinculoUsuarioCursoForm()

    vinculos = VinculoUsuarioCurso.objects.select_related('usuario', 'curso')
    return render(request, 'ppc/gerenciar_acessos.html', {
        'form': form,
        'vinculos': vinculos,
    })


@staff_member_required
@require_POST
def remover_acesso(request, vinculo_id):
    vinculo = get_object_or_404(VinculoUsuarioCurso, id=vinculo_id)
    vinculo.delete()
    messages.success(request, "Acesso removido.")
    return redirect('gerenciar_acessos')

def biblioteca(request):
    campos_busca = [
        {"id": "titulo", "rotulo": "Título"},
        {"id": "autores", "rotulo": "Autores"},
        {"id": "editora", "rotulo": "Editora"},
        {"id": "cidade", "rotulo": "Cidade"},
        {"id": "ano", "rotulo": "Ano"},
    ]

    campos_validos = {campo["id"] for campo in campos_busca}

    termo = request.GET.get("q", "").strip()
    busca_enviada = request.GET.get("buscar") == "1"

    if busca_enviada:
        campos_selecionados = [
            campo
            for campo in request.GET.getlist("campos")
            if campo in campos_validos
        ]
    else:
        # Na primeira abertura, todos os campos ficam selecionados.
        campos_selecionados = list(campos_validos)

    resultados = None
    erro_busca = None

    if busca_enviada:
        if not termo:
            erro_busca = "Digite um termo para pesquisar."

        elif not campos_selecionados:
            erro_busca = "Selecione pelo menos um campo para pesquisar."

        else:
            condicoes = []

            for campo in campos_selecionados:
                if campo == "ano":
                    # Para o ano, a pesquisa é exata.
                    if termo.isdecimal():
                        condicoes.append(Q(ano=int(termo)))
                else:
                    # Para textos, procura o termo em qualquer posição.
                    condicoes.append(
                        Q(**{f"{campo}__icontains": termo})
                    )

            if not condicoes:
                erro_busca = (
                    "Para pesquisar pelo campo Ano, "
                    "digite um ano numérico."
                )
            else:
                # Combina os campos selecionados usando OU.
                consulta = condicoes[0]

                for condicao in condicoes[1:]:
                    consulta |= condicao

                queryset = (
                    ReferenciaBibliografica.objects
                    .filter(consulta)
                    .order_by("titulo", "ano")
                )

                # Evita exibir milhares de referências de uma vez.
                paginador = Paginator(queryset, 25)
                resultados = paginador.get_page(
                    request.GET.get("page")
                )

    contexto = {
        "total_referenciabibliografica": (
            ReferenciaBibliografica.objects.count()
        ),
        "campos_busca": campos_busca,
        "campos_selecionados": campos_selecionados,
        "termo": termo,
        "busca_enviada": busca_enviada,
        "erro_busca": erro_busca,
        "resultados": resultados,
    }
    return render(request, "ppc/biblioteca.html", contexto)


def _buscar_bibliografia_vinculo(request, item, escopo):
    termo = request.GET.get("q", "").strip()
    tipo = request.GET.get("tipo", "basica")
    erro = None
    resultados = None
    if tipo not in {"basica", "complementar"}:
        tipo = "basica"
    if request.GET.get("buscar"):
        if len(termo) < 3:
            erro = "Digite ao menos 3 caracteres para pesquisar."
        else:
            qs = ReferenciaBibliografica.objects.filter(
                Q(titulo__icontains=termo) | Q(autores__icontains=termo)
            ).order_by("titulo", "ano", "id")
            resultados = Paginator(qs, 25).get_page(request.GET.get("page"))
    return render(request, "ppc/bibliografia/buscar_vinculo.html", {
        "item": item, "escopo": escopo, "termo": termo, "tipo": tipo,
        "erro": erro, "resultados": resultados,
        "busca_enviada": request.GET.get("buscar") == "1",
        "referencias_vinculadas": item.bibliografias.values_list("referencia_id", flat=True),
    })


@login_required
def buscar_bibliografia_ppc(request, componente_na_matriz_id):
    item = get_object_or_404(ComponenteNaMatriz, pk=componente_na_matriz_id)
    return _buscar_bibliografia_vinculo(request, item, "ppc")


@login_required
def buscar_bibliografia_referencia(request, item_id):
    item = get_object_or_404(ComponenteNaMatrizReferencia, pk=item_id)
    return _buscar_bibliografia_vinculo(request, item, "referencia")


@login_required
@require_POST
def vincular_bibliografia_ppc(request, componente_na_matriz_id):
    item = get_object_or_404(ComponenteNaMatriz, pk=componente_na_matriz_id)
    obra = get_object_or_404(ReferenciaBibliografica, pk=request.POST.get("referencia_id"))
    tipo = request.POST.get("tipo")
    if tipo not in {"basica", "complementar"}:
        return HttpResponseForbidden("Tipo bibliográfico inválido.")
    if item.bibliografias.filter(referencia=obra).exists():
        messages.info(request, "Esta obra já está vinculada ao componente.")
    else:
        manual = item.bibliografias.filter(
            referencia__isnull=True, tipo=tipo, titulo__iexact=obra.titulo,
            autores__iexact=obra.autores, ano=obra.ano,
        ).first()
        if manual:
            manual.referencia = obra
            manual.titulo, manual.autores = obra.titulo, obra.autores
            manual.editora, manual.cidade, manual.ano = obra.editora, obra.cidade, obra.ano
            manual.save()
            messages.success(request, "Referência manual correspondente vinculada à Biblioteca.")
        else:
            item.bibliografias.create(
                referencia=obra, tipo=tipo, titulo=obra.titulo, autores=obra.autores,
                editora=obra.editora, cidade=obra.cidade, ano=obra.ano,
            )
            messages.success(request, "Referência da Biblioteca vinculada ao componente.")
    return redirect("detalhe_componente_na_matriz", componente_na_matriz_id=item.pk)


@login_required
@require_POST
def vincular_bibliografia_referencia(request, item_id):
    item = get_object_or_404(ComponenteNaMatrizReferencia, pk=item_id)
    obra = get_object_or_404(ReferenciaBibliografica, pk=request.POST.get("referencia_id"))
    tipo = request.POST.get("tipo")
    if tipo not in {"basica", "complementar"}:
        return HttpResponseForbidden("Tipo bibliográfico inválido.")
    if item.bibliografias.filter(referencia=obra).exists():
        messages.info(request, "Esta obra já está vinculada ao componente.")
    else:
        manual = item.bibliografias.filter(
            referencia__isnull=True, tipo=tipo, titulo__iexact=obra.titulo,
            autores__iexact=obra.autores, ano=obra.ano,
        ).first()
        if manual:
            manual.referencia = obra
            manual.titulo, manual.autores = obra.titulo, obra.autores
            manual.editora, manual.cidade, manual.ano = obra.editora, obra.cidade, obra.ano
            manual.save()
            messages.success(request, "Referência manual correspondente vinculada à Biblioteca.")
        else:
            item.bibliografias.create(
                referencia=obra, tipo=tipo, titulo=obra.titulo, autores=obra.autores,
                editora=obra.editora, cidade=obra.cidade, ano=obra.ano,
            )
            messages.success(request, "Referência da Biblioteca vinculada ao componente.")
    return redirect("matriz_referencia_detalhe", matriz_id=item.matriz_id)


@login_required
@require_POST
def adicionar_bibliografia_manual_referencia(request, item_id):
    item = get_object_or_404(ComponenteNaMatrizReferencia, pk=item_id)
    form = BibliografiaReferenciaForm(request.POST)
    if form.is_valid():
        bib = form.save(commit=False)
        bib.componente_na_matriz_referencia = item
        bib.save()
        messages.success(request, "Referência manual adicionada; ela ficará pendente de vínculo à Biblioteca.")
    else:
        messages.error(request, "Não foi possível adicionar a referência manual. Confira os campos informados.")
    return redirect("matriz_referencia_detalhe", matriz_id=item.matriz_id)


@login_required
@require_POST
def excluir_bibliografia_referencia(request, bibliografia_referencia_id):
    bib = get_object_or_404(BibliografiaReferencia, pk=bibliografia_referencia_id)
    matriz_id = bib.componente_na_matriz_referencia.matriz_id
    bib.delete()
    messages.success(request, "Referência removida do componente.")
    return redirect("matriz_referencia_detalhe", matriz_id=matriz_id)

@login_required
def detalhe_componente_pendente(request, componente_id):
    componente = get_object_or_404(ComponenteCurricular, id=componente_id)
    return render(request, 'ppc/detalhe_pendente.html', {'componente': componente})

@login_required
def lista_equivalencias(request):
    grupos = list(
        GrupoEquivalenciaComponentes.objects.prefetch_related(
            'itens__item_matriz__componente', 'itens__item_matriz__matriz'
        ).order_by('-criado_em')
    )

    maior_grupo = max((g.itens.count() for g in grupos), default=1)
    colunas_cabecalho = range(1, max(maior_grupo, 1) + 1)

    for grupo in grupos:
        grupo.itens_ordenados = list(grupo.itens.all())  # já vem ordenado por 'ordem' via Meta.ordering
        grupo.celulas_vazias = range(maior_grupo - len(grupo.itens_ordenados))

    matrizes = MatrizReferenciaCurricular.objects.filter(matriz_anterior__isnull=False).select_related('curso').order_by('curso__nome', 'nome')
    return render(request, 'ppc/equivalencias/lista.html', {
        'grupos': grupos, 'colunas_cabecalho': colunas_cabecalho,
        'matrizes_com_edicao_anterior': matrizes,
    })


@login_required
def equivalencias_matriz(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, pk=matriz_id)
    dados = preparar_editor_equivalencias(matriz)
    if request.method == 'POST' and not dados['erro']:
        try:
            payload = json.loads(request.body) if request.content_type == "application/json" else {"pares": [v.split(":", 1) for v in request.POST.getlist("pares")]}
            pares = payload.get("pares", [])
            pares = [(int(p[0]), int(p[1])) for p in pares]
        except (ValueError, TypeError, KeyError, json.JSONDecodeError):
            if request.content_type == "application/json":
                return JsonResponse({"erro": "Seleção de equivalências inválida."}, status=400)
            messages.error(request, 'Seleção de equivalências inválida.')
            return render(request, 'ppc/equivalencias/matriz.html', dados)
        try:
            resumo = salvar_equivalencias_matriz(matriz.pk, pares, request.user)
        except ValidationError as erro:
            if request.content_type == "application/json":
                return JsonResponse({"erro": " ".join(erro.messages)}, status=400)
            messages.error(request, ' '.join(erro.messages))
        else:
            if request.content_type == "application/json":
                return JsonResponse({"ok": True, "resumo": resumo})
            messages.success(request, f"Equivalências salvas: {resumo['adicionadas']} adicionadas, {resumo['removidas']} removidas.")
            return redirect('equivalencias_matriz', matriz_id=matriz.pk)
    return render(request, 'ppc/equivalencias/matriz.html', dados)


@login_required
@require_GET
def candidatos_equivalencia_matriz(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, pk=matriz_id)
    try:
        item_id = int(request.GET.get("item_id", ""))
        anterior_id = int(request.GET.get("matriz_anterior_id", ""))
        offset = max(0, int(request.GET.get("offset", "0")))
        dados = candidatos_equivalencia_celula(
            matriz, item_id, anterior_id, request.GET.get("q", ""), offset,
        )
    except (ValueError, TypeError, ValidationError) as erro:
        mensagem = " ".join(erro.messages) if isinstance(erro, ValidationError) else "Parâmetros de busca inválidos."
        return JsonResponse({"erro": mensagem}, status=400)
    return JsonResponse(dados)


@login_required
@require_POST
def previa_equivalencias_matriz(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, pk=matriz_id)
    try:
        payload = json.loads(request.body)
        pares = payload.get("pares", [])
        if not isinstance(pares, list):
            raise ValueError
    except (json.JSONDecodeError, AttributeError, ValueError):
        return JsonResponse({"erro": "Seleção de equivalências inválida."}, status=400)
    dados = previsualizar_equivalencias_matriz(matriz, pares)
    html = render_to_string("ppc/equivalencias/_previa.html", dados, request=request)
    return JsonResponse({"html": html, "erro": dados.get("erro", "")})


@login_required
def equivalencias_ppc(request, ppc_id):
    ppc = get_object_or_404(PPC.objects.select_related("curso"), pk=ppc_id)
    dados = preparar_editor_ppc(ppc)
    if request.method == "POST":
        try:
            payload = json.loads(request.body)
            pares = [(int(p[0]), int(p[1])) for p in payload.get("pares", [])]
            if payload.get("preview"):
                from ppc.models import ComponenteNaMatriz
                from ppc.services.equivalencias_ppc import cadeia_ppc
                from ppc.services.equivalencias_matrizes import carga_equivalencia
                anteriores, erro = cadeia_ppc(ppc)
                antigos = {i.pk: i for i in ComponenteNaMatriz.objects.filter(ppc_id__in=[a.pk for a in anteriores]).select_related("componente", "ppc")}
                novos = {i.pk: i for i in ppc.matriz_componentes.select_related("componente")}
                rejeitados, validos, escolhidos = [], [], set()
                for novo_id, antigo_id in pares:
                    n, a = novos.get(novo_id), antigos.get(antigo_id)
                    motivo = "Componente não encontrado no PPC atual." if not n else "Componente anterior fora da linhagem." if not a else ""
                    if not motivo:
                        cn, ca = carga_equivalencia(n), carga_equivalencia(a)
                        if cn is None or ca is None: motivo = "Carga horária inválida."
                        elif ca < cn: motivo = f"CH anterior ({ca} h) menor que a CH nova ({cn} h)."
                    if motivo: rejeitados.append({"novo": novo_id, "antigo": antigo_id, "motivo": motivo})
                    else: escolhidos.add((novo_id, antigo_id))
                for n in sorted(novos.values(), key=lambda x: (x.periodo, x.componente.nome)):
                    celulas_preview = [{"ppc": ant, "pares": [antigos[a] for nn,a in escolhidos if nn == n.pk and antigos[a].ppc_id == ant.pk]} for ant in anteriores]
                    validos.append({"item": n, "celulas": celulas_preview, "has_equivalencia": any(c["pares"] for c in celulas_preview)})
                html = render_to_string("ppc/equivalencias/_previa_ppc.html", {"ppc": ppc, "ppcs_anteriores": anteriores, "linhas": validos, "rejeitados": rejeitados}, request=request)
                return JsonResponse({"html": html, "erro": erro})
            resumo = salvar_equivalencias_ppc(ppc.pk, pares, request.user)
            return JsonResponse({"ok": True, "resumo": resumo})
        except (ValueError, TypeError, KeyError, json.JSONDecodeError, ValidationError) as erro:
            msg = " ".join(erro.messages) if isinstance(erro, ValidationError) else "Seleção de equivalências inválida."
            return JsonResponse({"erro": msg}, status=400)
    dados["matriz"] = ppc
    return render(request, "ppc/equivalencias/ppc.html", dados)


@login_required
@require_GET
def candidatos_equivalencia_ppc(request, ppc_id):
    ppc = get_object_or_404(PPC, pk=ppc_id)
    try:
        data = candidatos_ppc(ppc, int(request.GET.get("item_id", "")), int(request.GET.get("matriz_anterior_id", "")), request.GET.get("q", ""), max(0, int(request.GET.get("offset", "0"))))
        return JsonResponse(data)
    except (ValueError, TypeError, ValidationError) as erro:
        return JsonResponse({"erro": " ".join(erro.messages) if isinstance(erro, ValidationError) else "Busca inválida."}, status=400)


@login_required
def criar_grupo_equivalencia(request):
    if request.method == 'POST':
        grupo = GrupoEquivalenciaComponentes.objects.create()
        return redirect('detalhe_grupo_equivalencia', grupo.id)
    return redirect('lista_equivalencias')


@login_required
def detalhe_grupo_equivalencia(request, grupo_id):
    grupo = get_object_or_404(GrupoEquivalenciaComponentes, id=grupo_id)
    itens = grupo.itens.select_related('item_matriz__componente', 'item_matriz__matriz')

    termo = request.GET.get('q', '').strip()
    resultados = []
    if termo:
        resultados = ComponenteNaMatrizReferencia.objects.select_related(
            'componente', 'matriz'
        ).filter(
            componente__nome__icontains=termo, componente__status='aprovado'
        ).exclude(
            id__in=itens.values_list('item_matriz_id', flat=True)
        )[:20]

    return render(request, 'ppc/equivalencias/detalhe.html', {
        'grupo': grupo, 'itens': itens, 'termo': termo, 'resultados': resultados,
    })


@login_required
@require_POST
def adicionar_item_equivalencia(request, grupo_id):
    grupo = get_object_or_404(GrupoEquivalenciaComponentes, id=grupo_id)
    item_matriz_id = request.POST.get('item_matriz_id')
    ordem = request.POST.get('ordem')
    item_matriz = get_object_or_404(ComponenteNaMatrizReferencia, id=item_matriz_id)

    item = ItemEquivalencia(grupo=grupo, item_matriz=item_matriz, ordem=ordem)
    try:
        item.full_clean()
    except ValidationError as erro:
        messages.error(request, " ".join(erro.messages))
    else:
        item.save()
    return redirect('detalhe_grupo_equivalencia', grupo.id)

@login_required
@require_POST
def remover_item_equivalencia(request, item_id):
    item = get_object_or_404(ItemEquivalencia, id=item_id)
    grupo_id = item.grupo_id
    item.delete()
    return redirect('detalhe_grupo_equivalencia', grupo_id)


@login_required
@require_POST
def definir_matriz_final(request, grupo_id):
    grupo = get_object_or_404(GrupoEquivalenciaComponentes, id=grupo_id)
    matriz_id = request.POST.get('matriz_id')
    grupo.matriz_final_id = matriz_id or None
    grupo.save(update_fields=['matriz_final'])
    return redirect('detalhe_grupo_equivalencia', grupo.id)

@login_required
def detalhe_componente_existente_matriz_referencia(request, matriz_id, componente_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, id=matriz_id)
    componente = get_object_or_404(ComponenteCurricular, id=componente_id)
    return render(request, "ppc/matrizes_referencia/detalhe_componente_existente.html", {
        "matriz": matriz,
        "componente": componente,
        "grupos_licenciatura": ComponenteNaMatrizReferencia.GRUPO_LICENCIATURA_CHOICES,
    })

@staff_member_required
def chamados_lista(request):
    chamados = Chamado.objects.select_related("usuario").all()
    return render(request, "ppc/chamados_lista.html", {"chamados": chamados})


@staff_member_required
@require_POST
def chamado_alternar_status(request, chamado_id):
    chamado = get_object_or_404(Chamado, id=chamado_id)
    if chamado.status == "pendente":
        chamado.status = "resolvido"
        chamado.resolvido_em = timezone.now()
    else:
        chamado.status = "pendente"
        chamado.resolvido_em = None
    chamado.save(update_fields=["status", "resolvido_em"])
    return redirect("chamados_lista")


@login_required
@require_POST
def abrir_chamado(request):
    mensagem = request.POST.get("mensagem", "").strip()
    tela = request.POST.get("tela", "")[:255]

    if not mensagem:
        return JsonResponse({"ok": False, "erro": "Mensagem vazia."}, status=400)

    chamado = Chamado.objects.create(
        usuario=request.user, tela=tela, mensagem=mensagem
    )

    try:
        send_mail(
            subject=f"[PPC-UFCAT] Novo chamado de {request.user}",
            message=f"Usuário: {request.user}\nTela: {tela}\n\nMensagem:\n{mensagem}",
            from_email=settings.DEFAULT_FROM_EMAIL,
            recipient_list=[settings.EMAIL_CHAMADOS_DESTINO],
            fail_silently=False,
        )
    except Exception as e:
        # o chamado já está salvo; o e-mail é best-effort, não deve quebrar a resposta
        print(f"[chamado] falha ao enviar e-mail: {e}")

    return JsonResponse({"ok": True, "id": chamado.id})

CAMPOS_SECOES_PPC = {
    "Apresentação": ["apresentacao_texto"],
    "Exposição de Motivos": ["exposicao_motivos"],
    "Objetivos": ["objetivo_geral", "objetivo_especifico"],
    "Princípios Norteadores": [
        "principios_geral", "principios_pratica_profissional",
        "principios_formacao_tecnica", "principios_formacao_etica_social",
        "principios_interdisciplinaridade", "principios_articulacao_teoria_pratica",
    ],
    "Expectativas da Formação": ["perfil_curso", "perfil_habilidades"],
    "TCC": ["tcc"],
    "Estágio": ["estagio"],
    "Atividades Complementares": ["atividades_complementares"],
    "Políticas Integradas": ["politicas_integrada"],
    "Avaliação do Ensino/Aprendizagem": ["avaliacao_ensino_aprendizagem"],
    "Avaliação do Projeto de Curso": ["avaliacao_projeto_curso"],
    "Qualificação de Docentes": ["qualificacao"],
    "Requisitos Legais": [
        "diretrizes_curriculares_nacionais_curso",
        "diretrizes_curriculares_nacionais_educacao_basica",
        "diretrizes_etnico_raciais_historia_cultura_afro_indigena",
        "diretrizes_educacao_direitos_humanos",
        "protecao_direitos_pessoa_transtorno_espectro_autista",
        "componente_curricular_libras",
        "politicas_educacao_ambiental",
        "diretrizes_formacao_professores_educacao_basica",
        "condicoes_acesso_pessoas_deficiencia_mobilidade_reduzida",
    ],
    "Bibliografias do PPC": ["bibliografias_ppc"],
    "Estrutura Curricular (descrição)": [
        "estrutura_curricular_descricao",
        "estrutura_curricular_informacoes_complementares",
    ],
}

CAMPOS_INFORMACOES_GERAIS = [
    "modalidade", "grau_academico", "turno_funcionamento", "carga_horaria_total",
    "numero_vagas_anuais", "duracao_minima_semestres", "duracao_media_semestres",
    "duracao_maxima_semestres", "diretor", "vice_diretor",
    "coordenador_curso", "vice_coordenador_curso",
]

CAMPOS_EAD = ["publico_alvo_ead", "ato_integracao_uab", "ato_credenciamento_mec", "polos_ead"]


@login_required
def duplicar_ppc(request, ppc_id):
    original = get_object_or_404(PPC, pk=ppc_id)

    if request.method == "POST":
        tipo_ppc = request.POST.get("tipo_ppc", "reformulacao")
        copiar_matriz = request.POST.get("copiar_matriz") == "on"
        secoes_marcadas = request.POST.getlist("secoes")

        relacoes_ignoradas = []
        periodos_acima_maximo = []
        with transaction.atomic():
            novo = PPC(curso=original.curso, tipo_ppc=tipo_ppc, status="rascunho")

            # Informações gerais sempre vêm como ponto de partida — são campos
            # obrigatórios do model, não dá pra criar o PPC sem eles preenchidos
            for campo in CAMPOS_INFORMACOES_GERAIS:
                setattr(novo, campo, getattr(original, campo))

            if original.modalidade == "ead":
                for campo in CAMPOS_EAD:
                    setattr(novo, campo, getattr(original, campo))

            for nome_secao, campos in CAMPOS_SECOES_PPC.items():
                if nome_secao in secoes_marcadas:
                    for campo in campos:
                        setattr(novo, campo, getattr(original, campo))

            novo.save()

            if copiar_matriz:
                # cópia em lote: o history grava tudo, mas não geramos um EventoEdicao por item
                with suspender_eventos_edicao():
                    mapa_componentes = {}  # id do ComponenteNaMatriz original -> novo ComponenteNaMatriz

                    for item in ComponenteNaMatriz.objects.filter(ppc=original):
                        novo_item = ComponenteNaMatriz(
                            ppc=novo, componente=item.componente,
                            periodo=item.periodo, natureza=item.natureza,
                            nucleo=item.nucleo, grupo_licenciatura=item.grupo_licenciatura,
                        )
                        novo_item.save(validar_limite_periodo=False)
                        if novo.duracao_maxima_semestres and item.periodo > novo.duracao_maxima_semestres:
                            periodos_acima_maximo.append({"componente": item.componente.nome, "periodo": item.periodo})
                        mapa_componentes[item.id] = novo_item

                        for bib in item.bibliografias.all():
                            Bibliografia.objects.create(
                                componente_na_matriz=novo_item,
                                referencia_id=bib.referencia_id,
                                tipo=bib.tipo, titulo=bib.titulo, autores=bib.autores,
                                editora=bib.editora, cidade=bib.cidade, ano=bib.ano,
                            )

                    # só depois que TODOS os componentes novos existem é que dá pra recriar
                    # as relações — uma relação exige as duas pontas já presentes
                    for relacao in RelacaoComponente.objects.filter(componente_na_matriz__ppc=original):
                        origem_nova = mapa_componentes.get(relacao.componente_na_matriz_id)
                        destino_novo = mapa_componentes.get(relacao.componente_relacionado_na_matriz_id)
                        if origem_nova and destino_novo:
                            existente = RelacaoComponente.objects.filter(
                                componente_na_matriz=origem_nova,
                                componente_relacionado_na_matriz=destino_novo,
                                tipo=relacao.tipo,
                            ).first()
                            if existente is None:
                                nova = RelacaoComponente(
                                    componente_na_matriz=origem_nova,
                                    componente_relacionado_na_matriz=destino_novo,
                                    tipo=relacao.tipo,
                                )
                                nova.save(validar_periodo=False)
                                mensagem = periodo_relacao_invalido(nova)
                                if mensagem:
                                    relacoes_ignoradas.append({"relacao": relacao, "mensagem": mensagem})

            if hasattr(original, "dinamica_ead") and original.modalidade == "ead":
                original_ead = original.dinamica_ead
                DinamicaEAD.objects.create(
                    ppc=novo,
                    dinamica_atividades_presenciais_distancia=original_ead.dinamica_atividades_presenciais_distancia,
                    recuperacao_estudos_permanencia=original_ead.recuperacao_estudos_permanencia,
                    componente_informatica_basica=original_ead.componente_informatica_basica,
                    atuacao_tutoria=original_ead.atuacao_tutoria,
                    atribuicoes_profissionais=original_ead.atribuicoes_profissionais,
                    material_didatico=original_ead.material_didatico,
                    ferramentas_comunicacao=original_ead.ferramentas_comunicacao,
                    carga_horaria_presencial_acompanhamento=original_ead.carga_horaria_presencial_acompanhamento,
                    armazenamento_gerenciamento_dados=original_ead.armazenamento_gerenciamento_dados,
                )
        registrar_evento(
            "ppc_duplicado", usuario=request.user, ppc=novo,
            origem_id=original.id, copiou_matriz=copiar_matriz, secoes=secoes_marcadas,
        )
        if relacoes_ignoradas:
            detalhes = " ".join(
                f"{r['relacao'].componente_na_matriz.componente.nome} → {r['relacao'].componente_relacionado_na_matriz.componente.nome}: {r['mensagem']}"
                for r in relacoes_ignoradas
            )
            messages.warning(request, f"{len(relacoes_ignoradas)} relação(ões) legada(s) inválida(s) foram copiadas. {detalhes}")
        if periodos_acima_maximo:
            detalhes = ", ".join(f"{item['componente']} (período {item['periodo']})" for item in periodos_acima_maximo)
            messages.warning(request, f"{len(periodos_acima_maximo)} componente(s) copiado(s) ficaram acima da duração máxima: {detalhes}.")

        return redirect("editar_apresentacao", ppc_id=novo.id)

    return render(request, "ppc/duplicar_ppc.html", {
        "original": original, "grupos": CAMPOS_SECOES_PPC,
    })


@login_required
def historico_ppc(request, ppc_id):
    ppc = get_object_or_404(PPC, pk=ppc_id)
    registros = ppc.history.all().order_by("-history_date")
    return render(request, "ppc/historico_ppc.html", {"ppc": ppc, "registros": registros})


@login_required
def reverter_ppc(request, ppc_id, history_id):
    ppc = get_object_or_404(PPC, pk=ppc_id)
    registro = get_object_or_404(ppc.history.all(), history_id=history_id)

    campos_diferentes = []
    for campo in ppc._meta.fields:
        if campo.name in ("id", "atualizado_em"):
            continue
        if getattr(ppc, campo.name, None) != getattr(registro, campo.name, None):
            campos_diferentes.append(str(campo.verbose_name or campo.name))

    if request.method == "POST":
        # .instance reconstrói um PPC com os valores daquele snapshot; salvar
        # por cima do atual gera uma NOVA entrada no histórico (reversão),
        # nunca apaga o que aconteceu depois — o "futuro" continua recuperável
        # revertendo de novo pra frente, se precisar.
        registro.instance.save()
        registrar_evento(
            "ppc_revertido", usuario=request.user, ppc=ppc,
            history_id=registro.history_id, campos=campos_diferentes,
        )
        return redirect("historico_ppc", ppc_id=ppc.id)

    return render(request, "ppc/confirmar_reversao.html", {
        "ppc": ppc, "registro": registro, "campos_diferentes": campos_diferentes,
    })

@login_required
def criar_matriz_referencia(request):
    if request.method == "POST":
        nome = request.POST.get("nome", "").strip()
        curso_id = request.POST.get("curso_id")
        try:
            numero_periodos = int(request.POST.get("numero_periodos", 8))
        except (TypeError, ValueError):
            numero_periodos = 8

        matriz = MatrizReferenciaCurricular.objects.create(
            nome=nome,
            curso=get_object_or_404(Curso, pk=curso_id) if curso_id else None,
            numero_periodos=numero_periodos,
        )
        return redirect("matriz_referencia_detalhe", matriz_id=matriz.id)

    return render(request, "ppc/matrizes_referencia/criar.html", {
        "cursos": Curso.objects.order_by("nome"),
    })


@login_required
@require_POST
def duplicar_matriz_referencia(request, matriz_id):
    original = get_object_or_404(MatrizReferenciaCurricular, pk=matriz_id)

    relacoes_ignoradas = []
    with transaction.atomic():
        copia = MatrizReferenciaCurricular.objects.create(
            curso=original.curso,
            nome=f"{original.nome or original.identificador_origem or 'Matriz'} (cópia)",
            numero_periodos=original.numero_periodos,
            grau_academico=original.grau_academico,
            matriz_anterior=original if request.POST.get("nova_edicao") == "1" else None,
        )
        total_itens = 0
        mapa_itens = {}
        # create() (e não bulk_create) para o history de cada item ser gravado;
        # o EventoEdicao por item fica suspenso e entra um único evento-resumo abaixo
        with suspender_eventos_edicao():
            for item in original.componentes.all():
                novo_item = ComponenteNaMatrizReferencia.objects.create(
                    matriz=copia, componente_id=item.componente_id,
                    periodo=item.periodo, natureza=item.natureza, ordem=item.ordem,
                    nucleo=item.nucleo, grupo_licenciatura=item.grupo_licenciatura,
                )
                mapa_itens[item.pk] = novo_item
                total_itens += 1
                for bib in item.bibliografias.all():
                    BibliografiaReferencia.objects.create(
                        componente_na_matriz_referencia=novo_item,
                        referencia_id=bib.referencia_id, tipo=bib.tipo,
                        titulo=bib.titulo, autores=bib.autores,
                        editora=bib.editora, cidade=bib.cidade, ano=bib.ano,
                    )
            relacoes_ignoradas = copiar_relacoes_entre_matrizes(original, copia)

    registrar_evento(
        "matriz_edicao_criada" if request.POST.get("nova_edicao") == "1" else "matriz_duplicada", usuario=request.user, curso=original.curso,
        origem_id=original.id, matriz_id=copia.id, itens=total_itens,
    )
    if relacoes_ignoradas:
        detalhes = " ".join(
            f"{r['relacao'].item.componente.nome} → {r['relacao'].item_relacionado.componente.nome}: {r['mensagem']}"
            for r in relacoes_ignoradas
        )
        messages.warning(request, f"{len(relacoes_ignoradas)} relação(ões) legada(s) inválida(s) foram copiadas. {detalhes}")
    return redirect("matriz_referencia_detalhe", matriz_id=copia.id)

@login_required
def buscar_componente_existente_matriz_referencia(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, pk=matriz_id)
    termo = request.GET.get("q", "").strip()
    resultados = []
    if termo:
        ja_usados = matriz.componentes.values_list("componente_id", flat=True)
        resultados = (
            ComponenteCurricular.objects
            .filter(nome__icontains=termo, status="aprovado")
            .exclude(id__in=ja_usados)
            .annotate(
                prioridade_ccu=Case(
                    When(ccu=True, then=Value(0)),
                    When(ccu=None, then=Value(1)),
                    When(ccu=False, then=Value(2)),
                    output_field=IntegerField(),
                )
            )
            .order_by('prioridade_ccu', 'nome')[:20]
        )
    return render(request, "ppc/matrizes_referencia/_resultados_busca.html", {
        "matriz": matriz, "resultados": resultados,
    })


@login_required
@require_POST
def adicionar_componente_existente_matriz_referencia(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, pk=matriz_id)
    componente = get_object_or_404(ComponenteCurricular, pk=request.POST.get("componente_id"))

    try:
        periodo = int(request.POST.get("periodo"))
    except (TypeError, ValueError):
        periodo = 1
    natureza = request.POST.get("natureza", "obrigatoria")
    grupo = request.POST.get("grupo_licenciatura") or None
    if grupo and grupo not in {str(valor) for valor, _ in ComponenteNaMatrizReferencia.GRUPO_LICENCIATURA_CHOICES}:
        return JsonResponse({"erro": "Grupo de licenciatura inválido."}, status=400)

    ComponenteNaMatrizReferencia.objects.get_or_create(
        matriz=matriz, componente=componente,
        defaults={"periodo": periodo, "natureza": natureza, "grupo_licenciatura": grupo},
    )
    return redirect("matriz_referencia_detalhe", matriz_id=matriz.id)


@login_required
def criar_componente_matriz_referencia(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, id=matriz_id)
    form = ComponenteMatrizReferenciaForm(request.POST or None, matriz=matriz)
    if request.method == 'POST' and form.is_valid():
        form.salvar_na_matriz(matriz, request.user)
        messages.success(request, "Componente criado e adicionado à matriz.")
        return redirect('matriz_referencia_detalhe', matriz_id=matriz.id)
    return render(request, 'ppc/matrizes_referencia/criar_componente.html', {
        'matriz': matriz,
        'form': form,
    })

@login_required
@require_POST
def remover_componente_matriz_referencia(request, item_id):
    item = get_object_or_404(ComponenteNaMatrizReferencia, pk=item_id)
    matriz_id = item.matriz_id
    item.delete()  # remove só o vínculo — o PROTECT protege o ComponenteCurricular do catálogo, não este link
    return redirect("matriz_referencia_detalhe", matriz_id=matriz_id)


@login_required
def historico_matriz_referencia(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, pk=matriz_id)
    registros = matriz.history.all().order_by("-history_date")
    return render(request, "ppc/matrizes_referencia/historico.html", {
        "matriz": matriz, "registros": registros,
    })


@login_required
def historico_item_matriz_referencia(request, item_id):
    item = get_object_or_404(ComponenteNaMatrizReferencia, pk=item_id)
    registros = item.history.all().order_by("-history_date")
    return render(request, "ppc/matrizes_referencia/historico_item.html", {
        "item": item, "registros": registros,
    })


@login_required
def selecionar_matriz_referencia_para_ppc(request, ppc_id):
    ppc = get_object_or_404(PPC, pk=ppc_id)
    matrizes = MatrizReferenciaCurricular.objects.filter(curso=ppc.curso)
    return render(request, "ppc/matrizes_referencia/selecionar_para_ppc.html", {
        "ppc": ppc, "matrizes": matrizes,
    })


def _montar_preview_aplicacao(ppc, matriz):
    """Categoriza cada item da matriz de referência frente ao que já existe
    no PPC. Nunca decide sozinho — só classifica, para o humano decidir."""
    itens_referencia = list(matriz.componentes.select_related("componente").prefetch_related(
        Prefetch("bibliografias", queryset=BibliografiaReferencia.objects.select_related("referencia").order_by("pk"))
    ))
    existentes = {
        cm.componente_id: cm
        for cm in ComponenteNaMatriz.objects.filter(ppc=ppc).select_related("componente").prefetch_related(
            Prefetch("bibliografias", queryset=Bibliografia.objects.select_related("referencia").order_by("pk"))
        )
    }

    preview = _categorizar_componentes(itens_referencia, existentes)
    preview["periodos_acima_maximo"] = [
        item for item in itens_referencia
        if ppc.duracao_maxima_semestres and item.periodo > ppc.duracao_maxima_semestres
    ]
    categorias = {
        item.pk: "novo" for item in preview["novos"]
    }
    categorias.update({item.pk: "igual" for item in preview["existentes_iguais"]})
    categorias.update({c["referencia"].pk: "conflito" for c in preview["conflitos"]})
    grupos, identicas = [], 0
    for item in itens_referencia:
        atual = existentes.get(item.componente_id)
        operacoes = calcular_operacoes_bibliografia(item, atual)
        categoria = categorias[item.pk]
        if not operacoes and categoria != "novo":
            identicas += 1
        if operacoes:
            for operacao in operacoes:
                operacao["item_id"] = item.pk
                operacao["componente_id"] = item.componente_id
            grupos.append({
                "item": item, "categoria": categoria,
                "operacoes": operacoes, "resumo": resumir_operacoes(operacoes),
                "componente_id": item.componente_id,
            })
    preview["bibliografias_diferencas"] = grupos
    preview["bibliografias_identicas"] = identicas
    preview["bibliografias_total_operacoes"] = sum(len(g["operacoes"]) for g in grupos)
    preview["graus_divergentes"] = ppc.grau_academico != matriz.grau_academico
    preview["grupos_licenciatura_divergentes"] = [
        {"referencia": item, "atual": existentes[item.componente_id]}
        for item in itens_referencia
        if ppc.grau_academico == "licenciatura"
        and item.componente_id in existentes
        and existentes[item.componente_id].grupo_licenciatura != item.grupo_licenciatura
    ]
    return preview


def _categorizar_componentes(itens_referencia, existentes):
    """Aplica a categorização usada no preview a uma coleção de componentes."""
    novos, existentes_iguais, conflitos = [], [], []

    for item in itens_referencia:
        existente = existentes.get(item.componente_id)
        if existente is None:
            novos.append(item)
        elif existente.periodo == item.periodo and existente.natureza == item.natureza:
            existentes_iguais.append(item)
        else:
            conflitos.append({"referencia": item, "atual": existente})

    return {"novos": novos, "existentes_iguais": existentes_iguais, "conflitos": conflitos}


def _executar_operacoes_bibliografia(matriz, operacoes, selecionadas, destinos, obsoletas=()):
    resumo = {"adicionadas": 0, "removidas": 0, "alteradas": 0, "substituidas": 0, "ignoradas": 0}
    ignoradas = []
    for chave in sorted(obsoletas):
        resumo["ignoradas"] += 1
        ignoradas.append(f"Operação {chave}: não existe mais no estado atual da matriz.")

    for op in operacoes:
        if op["key"] not in selecionadas:
            continue
        destino = destinos.get(op["item_id"])
        if destino is None:
            resumo["ignoradas"] += 1
            ignoradas.append(f"{op['rotulo']}: o componente não foi selecionado para aplicação.")
            continue
        atual = None
        if op["atual"] is not None:
            atual = Bibliografia.objects.filter(
                pk=op["atual"].pk, componente_na_matriz=destino,
            ).first()
            if atual is None:
                resumo["ignoradas"] += 1
                ignoradas.append(f"{op['rotulo']}: a bibliografia do PPC foi alterada ou removida após o preview.")
                continue
        fonte = None
        if op["fonte"] is not None:
            fonte = BibliografiaReferencia.objects.filter(
                pk=op["fonte"].pk,
                componente_na_matriz_referencia__matriz=matriz,
                componente_na_matriz_referencia__componente_id=op["componente_id"],
            ).first()
            if fonte is None:
                resumo["ignoradas"] += 1
                ignoradas.append(f"{op['rotulo']}: a referência foi alterada ou removida após o preview.")
                continue

        if op["acao"] == "remover":
            atual.delete()
            resumo["removidas"] += 1
            continue

        if op["acao"] in {"alterar_tipo", "substituir"}:
            atual.delete()
        Bibliografia.objects.create(
            componente_na_matriz=destino,
            referencia_id=fonte.referencia_id,
            tipo=fonte.tipo, titulo=fonte.titulo, autores=fonte.autores,
            editora=fonte.editora, cidade=fonte.cidade, ano=fonte.ano,
        )
        if op["acao"] == "adicionar":
            resumo["adicionadas"] += 1
        elif op["acao"] == "alterar_tipo":
            resumo["alteradas"] += 1
        else:
            resumo["substituidas"] += 1
    return resumo, ignoradas


@staff_member_required
def painel_institucional(request):

    pendencias_por_tipo = {
        escolha: 0 for escolha, _ in ImportacaoPendencia.TIPO_CHOICES
    }
    pendencias_por_tipo.update(
        ImportacaoPendencia.objects.filter(resolvido=False).values("tipo").annotate(total=Count("id")).values_list("tipo", "total")
    )

    cursos_sem_faixa = Curso.objects.filter(
        Q(carga_horaria_minima__isnull=True) | Q(carga_horaria_maxima__isnull=True)
    ).order_by("nome")
    ppcs_fora_da_faixa = PPC.objects.select_related("curso").filter(
        curso__carga_horaria_minima__isnull=False,
        curso__carga_horaria_maxima__isnull=False,
    ).filter(
        Q(carga_horaria_total__lt=F("curso__carga_horaria_minima"))
        | Q(carga_horaria_total__gt=F("curso__carga_horaria_maxima"))
    ).order_by("curso__nome", "id")

    resumo_bibliografias = resumo_pendencias_bibliograficas()
    total_relacoes_invalidas = contar_relacoes_invalidas()
    ppcs_sem_duracoes = PPC.objects.filter(
        duracao_minima_semestres=0, duracao_media_semestres=0, duracao_maxima_semestres=0
    )
    total_ppcs_fora_duracao = ppcs_fora_da_duracao_queryset().count()

    return render(request, "ppc/painel_institucional.html", {
        "total_matrizes": MatrizReferenciaCurricular.objects.count(),
        "matrizes_sem_curso": MatrizReferenciaCurricular.objects.filter(curso__isnull=True).count(),
        "total_pendencias": ImportacaoPendencia.objects.filter(resolvido=False).count(),
        "pendencias_por_tipo": [
            {"codigo": codigo, "rotulo": rotulo, "total": pendencias_por_tipo[codigo]}
            for codigo, rotulo in ImportacaoPendencia.TIPO_CHOICES
        ],
        "cursos_sem_faixa": cursos_sem_faixa,
        "ppcs_fora_da_faixa": ppcs_fora_da_faixa,
        "componentes_pendentes": ComponenteCurricular.objects.filter(status="pendente").count(),
        "resumo_bibliografias": resumo_bibliografias,
        "total_relacoes_invalidas": total_relacoes_invalidas,
        "ppcs_sem_duracoes_count": ppcs_sem_duracoes.count(),
        "ppcs_fora_duracao_count": total_ppcs_fora_duracao,
    })


@login_required
def detalhe_pendencias_bibliografia_ppc(request, ppc_id):
    ppc = get_object_or_404(PPC.objects.select_related("curso"), pk=ppc_id)
    pagina = Paginator(queryset_componentes_pendentes_ppc(ppc.pk).order_by("periodo", "componente__nome"), 25).get_page(request.GET.get("page"))
    return render(request, "ppc/painel/detalhe_pendencias_bibliografia.html", {
        "escopo": "PPC", "entidade": ppc, "curso": ppc.curso,
        "itens": pagina, "pagina": pagina,
    })


@login_required
def detalhe_pendencias_bibliografia_matriz(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular.objects.select_related("curso"), pk=matriz_id)
    pagina = Paginator(queryset_componentes_pendentes_matriz(matriz.pk).order_by("periodo", "componente__nome"), 25).get_page(request.GET.get("page"))
    return render(request, "ppc/painel/detalhe_pendencias_bibliografia.html", {
        "escopo": "Matriz de referência", "entidade": matriz, "curso": matriz.curso,
        "itens": pagina, "pagina": pagina,
    })


@login_required
def detalhe_relacoes_invalidas(request):
    pagina = Paginator(relacoes_invalidas(), 25).get_page(request.GET.get("page"))
    return render(request, "ppc/painel/detalhe_relacoes_invalidas.html", {
        "relacoes": pagina, "pagina": pagina,
    })


@_bloquear_nde_em_detalhes_staff
@staff_member_required
def detalhe_ppcs_sem_duracoes(request):
    queryset = PPC.objects.select_related("curso").filter(
        duracao_minima_semestres=0, duracao_media_semestres=0, duracao_maxima_semestres=0
    ).order_by("curso__nome", "pk")
    pagina = Paginator(queryset, 25).get_page(request.GET.get("page"))
    return render(request, "ppc/painel/detalhe_ppcs_sem_duracoes.html", {
        "ppcs": pagina, "pagina": pagina,
    })


@_bloquear_nde_em_detalhes_staff
@staff_member_required
def detalhe_ppcs_fora_duracao(request):
    queryset = ppcs_fora_da_duracao_queryset().select_related("curso").order_by("curso__nome", "pk")
    pagina = Paginator(queryset, 25).get_page(request.GET.get("page"))
    ids_ppcs = [ppc.pk for ppc in pagina]
    componentes_fora = {ppc_id: [] for ppc_id in ids_ppcs}
    for item in ComponenteNaMatriz.objects.filter(ppc_id__in=ids_ppcs).select_related("componente", "ppc"):
        if item.ppc.duracao_maxima_semestres and item.periodo > item.ppc.duracao_maxima_semestres:
            componentes_fora[item.ppc_id].append(item)
    linhas = [{"ppc": ppc, "componentes_fora": componentes_fora[ppc.pk]} for ppc in pagina]
    return render(request, "ppc/painel/detalhe_ppcs_fora_duracao.html", {
        "linhas": linhas, "pagina": pagina,
    })


@login_required
def comparar_ppcs(request, ppc_id):
    ppc = get_object_or_404(PPC.objects.select_related("curso"), pk=ppc_id)
    ppcs_do_curso = PPC.objects.filter(curso=ppc.curso).exclude(pk=ppc.pk).order_by("criado_em", "id")
    ppc_comparado = None
    comparacao = None

    if ppc_comparado_id := request.GET.get("ppc_comparado"):
        ppc_comparado = get_object_or_404(
            PPC.objects.select_related("curso"), pk=ppc_comparado_id, curso=ppc.curso
        )
        itens_comparados = ppc_comparado.matriz_componentes.select_related("componente")
        itens_atuais = {
            item.componente_id: item
            for item in ppc.matriz_componentes.select_related("componente")
        }
        comparacao = _categorizar_componentes(itens_comparados, itens_atuais)

    return render(request, "ppc/comparar_ppcs.html", {
        "ppc": ppc,
        "ppcs_do_curso": ppcs_do_curso,
        "ppc_comparado": ppc_comparado,
        "comparacao": comparacao,
        "diferenca_carga_horaria": (
            ppc_comparado.carga_horaria_total - ppc.carga_horaria_total
            if ppc_comparado else None
        ),
    })


@login_required
def preview_aplicar_matriz_referencia(request, ppc_id, matriz_id):
    ppc = get_object_or_404(PPC, pk=ppc_id)
    matriz = get_object_or_404(MatrizReferenciaCurricular, pk=matriz_id)
    preview = _montar_preview_aplicacao(ppc, matriz)

    if request.method == "POST":
        with transaction.atomic(), suspender_eventos_edicao():
            # Recalcula dentro da transação; nenhum dado do formulário é usado
            # como fonte de bibliografias ou como decisão de diferença.
            preview = _montar_preview_aplicacao(ppc, matriz)
            novos_selecionados = [
                item for item in preview["novos"]
                if request.POST.get(f"aplicar_novo_{item.id}") == "on"
            ]
            conflitos_selecionados = [
                conflito for conflito in preview["conflitos"]
                if request.POST.get(f"conflito_{conflito['referencia'].id}") == "usar_referencia"
            ]
            ids_elegiveis = {item.pk for item in preview["existentes_iguais"]}
            ids_elegiveis.update(item.pk for item in novos_selecionados)
            ids_elegiveis.update(c["referencia"].pk for c in conflitos_selecionados)
            grupos_elegiveis = [g for g in preview["bibliografias_diferencas"] if g["item"].pk in ids_elegiveis]
            operacoes_elegiveis = [op for grupo in grupos_elegiveis for op in grupo["operacoes"]]
            selecionadas = set(request.POST.getlist("bibliografia_operacao"))
            if preview["bibliografias_total_operacoes"] and request.POST.get("confirmar_bibliografias") != "on":
                messages.error(request, "Confira as diferenças bibliográficas e confirme as operações antes de aplicar a matriz.")
                return render(request, "ppc/matrizes_referencia/preview_aplicacao.html", {
                    "ppc": ppc, "matriz": matriz, "preview": preview,
                })

            periodos_acima_maximo = [
                item for item in novos_selecionados + [c["referencia"] for c in conflitos_selecionados]
                if ppc.duracao_maxima_semestres and item.periodo > ppc.duracao_maxima_semestres
            ]
            componentes_nao_atualizados = []
            aplicados = 0
            vinculos_destino = {}
            for item in preview["existentes_iguais"]:
                vinculos_destino[item.pk] = ComponenteNaMatriz.objects.filter(
                    ppc=ppc, componente_id=item.componente_id
                ).first()

            for item in novos_selecionados:
                vinculo = ComponenteNaMatriz(
                    ppc=ppc, componente=item.componente,
                    periodo=item.periodo, natureza=item.natureza,
                    nucleo=item.nucleo,
                    grupo_licenciatura=(item.grupo_licenciatura if ppc.grau_academico == "licenciatura" else None),
                )
                vinculo.save(validar_limite_periodo=False)
                vinculos_destino[item.pk] = vinculo
                aplicados += 1

            for conflito in preview["conflitos"]:
                if conflito in conflitos_selecionados:
                    atual = conflito["atual"]
                    periodo_anterior = atual.periodo
                    natureza_anterior = atual.natureza
                    atual.periodo = conflito["referencia"].periodo
                    atual.natureza = conflito["referencia"].natureza
                    try:
                        atual.save(update_fields=["periodo", "natureza"], validar_limite_periodo=False)
                    except ValidationError as exc:
                        atual.periodo, atual.natureza = periodo_anterior, natureza_anterior
                        componentes_nao_atualizados.append({
                            "componente": atual.componente.nome,
                            "mensagem": "; ".join(exc.messages),
                        })
                    else:
                        aplicados += 1
                    vinculos_destino[conflito["referencia"].pk] = atual
                # "manter_ppc" ou nenhuma escolha -> não mexe, PPC continua como está

            todas_chaves_atuais = {
                op["key"] for grupo in preview["bibliografias_diferencas"] for op in grupo["operacoes"]
            }
            chaves_elegiveis = {op["key"] for op in operacoes_elegiveis}
            resumo_bibliografias, operacoes_ignoradas = _executar_operacoes_bibliografia(
                matriz, operacoes_elegiveis, selecionadas & chaves_elegiveis,
                vinculos_destino, obsoletas=selecionadas - todas_chaves_atuais,
            )

            resultado_relacoes = copiar_relacoes_referencia_para_ppc(matriz, ppc)

        registrar_evento(
            "matriz_aplicada", usuario=request.user, ppc=ppc,
            matriz_referencia_id=matriz.pk, componentes_aplicados=aplicados,
            bibliografias_adicionadas=resumo_bibliografias["adicionadas"],
            bibliografias_removidas=resumo_bibliografias["removidas"],
            bibliografias_alteradas=resumo_bibliografias["alteradas"],
            bibliografias_substituidas=resumo_bibliografias["substituidas"],
            bibliografias_ignoradas=resumo_bibliografias["ignoradas"],
            relacoes_copiadas=resultado_relacoes["criadas"],
            relacoes_ignoradas=len(resultado_relacoes["invalidas"]),
        )

        return render(request, "ppc/matrizes_referencia/aplicacao_concluida.html", {
            "ppc": ppc, "matriz": matriz, "aplicados": aplicados,
            "relacoes_ignoradas": resultado_relacoes["invalidas"],
            "periodos_acima_maximo": periodos_acima_maximo,
            "componentes_nao_atualizados": componentes_nao_atualizados,
            "resumo_bibliografias": resumo_bibliografias,
            "operacoes_bibliograficas_ignoradas": operacoes_ignoradas,
        })

    return render(request, "ppc/matrizes_referencia/preview_aplicacao.html", {
        "ppc": ppc, "matriz": matriz, "preview": preview,
    })


@login_required
def lista_pendencias_por_matriz(request):
    grupos = (
        ImportacaoPendencia.objects
        .filter(tipo="ambigua", resolvido=False)
        .values("identificador_origem")
        .annotate(total=Count("id"))
        .order_by("-total")
    )
    return render(request, "ppc/pendencias/lista_por_matriz.html", {"grupos": grupos})


@login_required
def revisar_pendencias_matriz(request, identificador_origem):
    pendencias = (
        ImportacaoPendencia.objects
        .filter(tipo="ambigua", resolvido=False, identificador_origem=identificador_origem)
        .order_by("linha_origem")
    )

    if request.method == "POST":
        resolvidas = 0
        for pendencia in pendencias:
            decisao = request.POST.get(f"decisao_{pendencia.id}")
            if decisao in ("obrigatoria", "optativa"):
                if resolver_pendencia_ambigua(pendencia, decisao):
                    resolvidas += 1
            # decisao == "pendente" ou ausente -> não faz nada, continua pendente

        return render(
            request, "ppc/pendencias/revisao_concluida.html",
            {"identificador_origem": identificador_origem, "resolvidas": resolvidas},
        )

    return render(
        request, "ppc/pendencias/revisar_por_matriz.html",
        {"identificador_origem": identificador_origem, "pendencias": pendencias},
    )

@login_required
def lista_matrizes_referencia(request):
    matrizes = (
        MatrizReferenciaCurricular.objects
        .select_related("curso")
        .order_by("curso__nome", "nome")
    )
    return render(request, "ppc/matrizes_referencia/lista.html", {"matrizes": matrizes})


@login_required
def matriz_referencia_detalhe(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, pk=matriz_id)
    return render(request, "ppc/matrizes_referencia/detalhe.html", _contexto_detalhe_matriz(matriz))


def _contexto_detalhe_matriz(matriz):
    itens = list(
        matriz.componentes
        .select_related("componente")
        .prefetch_related("bibliografias")
        .order_by("periodo", "ordem")
    )
    pendencias = resumo_pendencias_matriz(matriz)
    horas = calcular_horas_semanais(itens)
    for item in itens:
        item.pendencia_bibliografia = pendencias["pendencias_por_item"].get(item.pk, "")
        item.horas_semanais = horas[item.pk]
        item.carga_total_segura = _carga_total_legada_segura(item.componente)

    periodos = {n: [] for n in range(1, matriz.numero_periodos + 1)}
    for item in itens:
        # se algum item tiver período fora do range configurado (ex: veio do
        # import com período 9 mas numero_periodos ficou em 8), a coluna dele
        # aparece mesmo assim — não queremos esconder dado existente
        periodos.setdefault(item.periodo, []).append(item)

    context = {
        "matriz": matriz,
        "periodos": sorted(periodos.items()),
        "cursos": Curso.objects.order_by("nome"),
        "matrizes_do_mesmo_curso": MatrizReferenciaCurricular.objects.filter(curso_id=matriz.curso_id).exclude(pk=matriz.pk).order_by("nome") if matriz.curso_id else MatrizReferenciaCurricular.objects.none(),
        "resumo_pendencias_matriz": pendencias,
        "detalhe_pendencias_url": reverse("detalhe_pendencias_bibliografia_matriz", args=[matriz.pk]),
        "resumo_periodos": _resumo_periodos_matriz(periodos),
        "tem_acex": any((item.componente.carga_horaria_acex or 0) for item in itens),
    }
    return context


def _resumo_periodos_matriz(periodos):
    resumo = []
    for periodo, itens in sorted(periodos.items()):
        obrigatorias = [i for i in itens if i.natureza == "obrigatoria"]
        optativas = [i for i in itens if i.natureza == "optativa"]
        semanal = calcular_horas_semanais(itens)
        resumo.append({
            "periodo": periodo, "componentes": len(itens),
            "obrigatorias": len(obrigatorias), "optativas": len(optativas),
            "obrig_h_sem": sum(semanal[i.pk]["total"] for i in obrigatorias),
            "opt_h_sem": sum(semanal[i.pk]["total"] for i in optativas),
            "teorica": sum(i.componente.carga_horaria_teorica or 0 for i in itens if i.componente.tipo != "atividade"),
            "pratica": sum(i.componente.carga_horaria_pratica or 0 for i in itens if i.componente.tipo != "atividade"),
            "ch_total": sum(_carga_total_legada_segura(i.componente) for i in itens),
            "acex": sum(i.componente.carga_horaria_acex or 0 for i in itens),
            "obrig_h_sem_texto": formatar_horas_semanais_valor(sum(semanal[i.pk]["total"] for i in obrigatorias)),
            "opt_h_sem_texto": formatar_horas_semanais_valor(sum(semanal[i.pk]["total"] for i in optativas)),
        })
    return resumo


def _carga_total_legada_segura(componente):
    try:
        return componente.carga_horaria_computada_total or 0
    except (TypeError, ValueError):
        return 0


@login_required
@require_POST
def editar_item_matriz_referencia(request, item_id):
    item = get_object_or_404(ComponenteNaMatrizReferencia.objects.select_related("matriz", "componente"), pk=item_id)
    try:
        dados = json.loads(request.body)
        if not isinstance(dados, dict):
            raise ValueError("O corpo da requisição deve ser um objeto JSON.")
        # Campos alheios ao vínculo (inclusive catálogo) são ignorados.
        grupo = (dados.get("grupo_licenciatura") or None) if "grupo_licenciatura" in dados else item.grupo_licenciatura
        if item.matriz.grau_academico != "licenciatura":
            grupo = None
        if grupo and str(grupo) not in {str(v) for v, _ in ComponenteNaMatrizReferencia.GRUPO_LICENCIATURA_CHOICES}:
            raise ValueError("Grupo de licenciatura inválido.")
        item.nucleo = dados.get("nucleo", item.nucleo)
        item.natureza = dados.get("natureza", item.natureza)
        item.grupo_licenciatura = grupo
        item.full_clean()
        with transaction.atomic():
            item.save()
    except (KeyError, TypeError, ValueError, json.JSONDecodeError, ValidationError) as exc:
        erro = "; ".join(exc.messages) if isinstance(exc, ValidationError) else str(exc)
        return JsonResponse({"erro": erro or "Dados inválidos."}, status=400)
    semanal = calcular_horas_semanais([item])[item.pk]
    return JsonResponse({
        "ok": True, "natureza": item.natureza,
        "nucleo": item.nucleo, "grupo_licenciatura": item.grupo_licenciatura,
        "horas_semanais": semanal["total"], "horas_semanais_texto": semanal["texto"],
        "horas_semanais_detalhe": semanal["detalhe"],
        "horas_semanais_teoricas": semanal["teorica"],
        "horas_semanais_praticas": semanal["pratica"],
        "atividade": semanal["atividade"],
    })


@login_required
def exportar_matriz_referencia_xlsx(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, pk=matriz_id)
    if openpyxl is None:
        messages.error(request, "A exportação Excel está indisponível neste ambiente porque o componente openpyxl não está instalado.")
        return redirect("matriz_referencia_detalhe", matriz_id=matriz.pk)
    itens = list(matriz.componentes.select_related("componente").prefetch_related("bibliografias").order_by("periodo", "ordem", "componente__nome"))
    colunas = [
        "matriz_curricular", "codigo", "id_componente_interno", "nucleo", "disciplina", "fase",
        "cargahorariateorica", "cargahorariapratica", "cargahorariaacex", "ementa",
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matriz"
    ws.append(colunas)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24527A")
    identificador = matriz.identificador_origem or f"EXPORT-{matriz.pk}"
    for item in itens:
        componente = item.componente
        nucleo = (
            "NÚCLEO COMUM" if item.nucleo == "NC" else
            "NÚCLEO ESPECÍFICO OPTATIVO" if item.natureza == "optativa" else
            "NÚCLEO ESPECÍFICO OBRIGATÓRIO"
        )
        ws.append([
            identificador, componente.codigo, item.pk, nucleo,
            componente.nome, item.periodo, componente.carga_horaria_teorica,
            componente.carga_horaria_pratica, componente.carga_horaria_acex or 0,
            componente.ementa,
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in {"A": 24, "B": 16, "C": 22, "D": 32, "E": 42, "F": 12, "G": 20, "H": 20, "I": 18, "J": 70}.items():
        ws.column_dimensions[col].width = width

    # O diagrama é a primeira aba aberta e reproduz o layout que o Apps Script exporta.
    diagrama = wb.create_sheet("Diagrama")
    fases = sorted({item.periodo for item in itens})
    largura_bloco = 3
    total_colunas = max(1, len(fases) * largura_bloco)
    titulo_fim = openpyxl.utils.get_column_letter(total_colunas)
    diagrama.merge_cells(f"A1:{titulo_fim}1")
    diagrama["A1"] = f"MATRIZ CURRICULAR — {matriz}"
    diagrama["A1"].font = Font(bold=True, size=15)
    diagrama["A1"].fill = PatternFill("solid", fgColor="DCEAF7")
    diagrama["A1"].alignment = Alignment(vertical="center")
    diagrama.row_dimensions[1].height = 28
    cores = {
        "NC": "FFF2CC",
        "NE_OBRIGATORIA": "E4F5E8",
        "NE_OPTATIVA": "E9E8FF",
        "ACEX": "FCE8E8",
    }
    borda = Side(style="thin", color="A9B7C7")
    for indice, fase in enumerate(fases):
        primeira_coluna = indice * largura_bloco + 1
        ultima_coluna = primeira_coluna + largura_bloco - 1
        inicio = openpyxl.utils.get_column_letter(primeira_coluna)
        fim = openpyxl.utils.get_column_letter(ultima_coluna)
        for col in range(primeira_coluna, ultima_coluna + 1):
            diagrama.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 16
        diagrama.merge_cells(start_row=3, start_column=primeira_coluna, end_row=3, end_column=ultima_coluna)
        celula_fase = diagrama.cell(3, primeira_coluna, f"{fase}º PERÍODO")
        celula_fase.font = Font(bold=True)
        celula_fase.fill = PatternFill("solid", fgColor="DDE7F2")
        celula_fase.alignment = Alignment(horizontal="center", vertical="center")
        linha = 5
        por_fase = [item for item in itens if item.periodo == fase]
        for item in por_fase:
            componente = item.componente
            nucleo = "Núcleo Comum" if item.nucleo == "NC" else (
                "Núcleo Específico Optativo" if item.natureza == "optativa"
                else "Núcleo Específico Obrigatório"
            )
            cor = "ACEX" if (componente.carga_horaria_acex or 0) > 0 else (
                "NC" if item.nucleo == "NC" else
                "NE_OPTATIVA" if item.natureza == "optativa" else "NE_OBRIGATORIA"
            )
            fundo = PatternFill("solid", fgColor=cores[cor])
            estilo_borda = Border(left=borda, right=borda, top=borda, bottom=borda)
            diagrama.merge_cells(start_row=linha, start_column=primeira_coluna, end_row=linha, end_column=ultima_coluna)
            diagrama.cell(linha, primeira_coluna, componente.codigo or "Sem código").font = Font(bold=True)
            diagrama.merge_cells(start_row=linha + 1, start_column=primeira_coluna, end_row=linha + 2, end_column=ultima_coluna)
            nome_celula = diagrama.cell(linha + 1, primeira_coluna, componente.nome)
            nome_celula.font = Font(bold=True)
            nome_celula.alignment = Alignment(vertical="center", wrap_text=True)
            diagrama.merge_cells(start_row=linha + 3, start_column=primeira_coluna, end_row=linha + 3, end_column=ultima_coluna)
            diagrama.cell(linha + 3, primeira_coluna, nucleo)
            ch_total = componente.carga_horaria_computada_total
            resumo = f"{componente.carga_horaria_teorica}T · {componente.carga_horaria_pratica}P · {componente.carga_horaria_acex or 0} ACEx · {ch_total} Total"
            diagrama.merge_cells(start_row=linha + 4, start_column=primeira_coluna, end_row=linha + 4, end_column=ultima_coluna)
            diagrama.cell(linha + 4, primeira_coluna, resumo)
            for r in range(linha, linha + 5):
                for c in range(primeira_coluna, ultima_coluna + 1):
                    cell = diagrama.cell(r, c)
                    cell.fill = fundo
                    cell.border = estilo_borda
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
            diagrama.row_dimensions[linha + 1].height = 27
            diagrama.row_dimensions[linha + 2].height = 27
            linha += 7
    diagrama.freeze_panes = "A4"

    bibliografias = wb.create_sheet("Bibliografias")
    bibliografias.append(["fase", "codigo", "componente", "tipo", "titulo", "autores", "editora", "cidade", "ano", "id_referencia_biblioteca"])
    for item in itens:
        for bib in item.bibliografias.all():
            bibliografias.append([
                item.periodo, item.componente.codigo, item.componente.nome, bib.tipo,
                bib.titulo, bib.autores, bib.editora, bib.cidade, bib.ano,
                bib.referencia_id or "",
            ])
    for cell in bibliografias[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24527A")
    bibliografias.freeze_panes = "A2"
    bibliografias.auto_filter.ref = bibliografias.dimensions

    relacoes = wb.create_sheet("Relacoes")
    relacoes.append(["codigo_componente", "componente", "tipo", "codigo_relacionado", "componente_relacionado", "periodo_componente", "periodo_relacionado"])
    for relacao in RelacaoComponenteReferencia.objects.filter(item__matriz=matriz).select_related(
        "item__componente", "item_relacionado__componente"
    ):
        relacoes.append([
            relacao.item.componente.codigo, relacao.item.componente.nome, relacao.tipo,
            relacao.item_relacionado.componente.codigo, relacao.item_relacionado.componente.nome,
            relacao.item.periodo, relacao.item_relacionado.periodo,
        ])
    for cell in relacoes[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24527A")
    relacoes.freeze_panes = "A2"
    relacoes.auto_filter.ref = relacoes.dimensions

    wb.active = wb.sheetnames.index("Diagrama")
    buffer = BytesIO()
    wb.save(buffer)
    nome = slugify(matriz.nome or matriz.identificador_origem or f"matriz-{matriz.pk}") or f"matriz-{matriz.pk}"
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="matriz-referencia-{nome}.xlsx"'
    return response


@login_required
@require_POST
def vincular_curso_matriz_referencia(request, matriz_id):
    matriz = get_object_or_404(MatrizReferenciaCurricular, pk=matriz_id)

    curso_id = request.POST.get("curso_id")
    nome = request.POST.get("nome", "").strip()

    try:
        numero_periodos = int(request.POST.get("numero_periodos", matriz.numero_periodos))
        if numero_periodos < 1:
            raise ValueError
    except (TypeError, ValueError):
        numero_periodos = matriz.numero_periodos  # ignora valor inválido, mantém o que já tinha

    grau = request.POST.get("grau_academico") or None
    matriz_anterior_id = request.POST.get("matriz_anterior_id") or None
    matriz_anterior = None
    if matriz_anterior_id:
        matriz_anterior = get_object_or_404(MatrizReferenciaCurricular, pk=matriz_anterior_id)
        curso_destino_id = curso_id or None
        if not curso_destino_id or matriz_anterior.curso_id != int(curso_destino_id):
            messages.error(request, "A matriz anterior deve estar vinculada ao mesmo curso.")
            return redirect("matriz_referencia_detalhe", matriz_id=matriz.pk)
    if grau not in {valor for valor, _ in PPC.GRAU_CHOICES} and grau is not None:
        grau = matriz.grau_academico
    grupos_para_limpar = 0
    if matriz.grau_academico == "licenciatura" and grau != "licenciatura":
        grupos_para_limpar = matriz.componentes.filter(grupo_licenciatura__isnull=False).count()
        if grupos_para_limpar and request.POST.get("confirmar_limpeza_grupos") != "1":
            return render(request, "ppc/matrizes_referencia/detalhe.html", {
                **_contexto_detalhe_matriz(matriz),
                "confirmacao_limpeza_grupos": {
                    "quantidade": grupos_para_limpar,
                    "nome": nome,
                    "curso_id": curso_id or "",
                    "numero_periodos": numero_periodos,
                    "grau_academico": grau or "",
                    "matriz_anterior_id": matriz_anterior.pk if matriz_anterior else "",
                },
            })

    with transaction.atomic():
        matriz.curso = get_object_or_404(Curso, pk=curso_id) if curso_id else None
        matriz.nome = nome
        matriz.numero_periodos = numero_periodos
        matriz.grau_academico = grau
        matriz.matriz_anterior = matriz_anterior
        try:
            matriz.save()
        except ValidationError as erro:
            messages.error(request, " ".join(erro.messages))
            return redirect("matriz_referencia_detalhe", matriz_id=matriz.pk)
        if grupos_para_limpar:
            with suspender_eventos_edicao():
                for item in matriz.componentes.filter(grupo_licenciatura__isnull=False):
                    item.grupo_licenciatura = None
                    item.save(update_fields=["grupo_licenciatura"])
            registrar_evento(
                "grupos_licenciatura_matriz_limpos", usuario=request.user,
                curso=matriz.curso, matriz_id=matriz.pk, grupos=grupos_para_limpar,
            )
    return redirect("matriz_referencia_detalhe", matriz_id=matriz.id)


@login_required
@require_POST
def mover_item_matriz_referencia(request, item_id):
    """Recebe o novo período/posição de um item arrastado e reordena os
    irmãos do período de destino para não haver 'ordem' duplicada."""
    try:
        payload = json.loads(request.body)
        novo_periodo = int(payload["novo_periodo"])
        nova_ordem = int(payload["nova_ordem"])
    except (KeyError, ValueError, TypeError, json.JSONDecodeError):
        return JsonResponse({"erro": "Dados inválidos."}, status=400)

    item = get_object_or_404(ComponenteNaMatrizReferencia, pk=item_id)

    try:
        with transaction.atomic():
            item.periodo = novo_periodo
            item.ordem = nova_ordem
            item.save(update_fields=["periodo", "ordem"])

            irmaos = (
                ComponenteNaMatrizReferencia.objects
                .filter(matriz=item.matriz, periodo=novo_periodo)
                .exclude(pk=item.pk)
                .order_by("ordem")
            )
            posicao = 0
            for irmao in irmaos:
                if posicao == nova_ordem:
                    posicao += 1
                if irmao.ordem != posicao:
                    irmao.ordem = posicao
                    irmao.save(update_fields=["periodo", "ordem"])
                posicao += 1
    except ValidationError as exc:
        mensagem = "; ".join(exc.messages)
        return JsonResponse({"erro": mensagem}, status=400)

    return JsonResponse({"ok": True})


CAMPOS_PPC_VALIDOS = {f.name for f in PPC._meta.get_fields()}

@staff_member_required(login_url='login')
def fila_aprovacao_componentes(request):
    pendentes = ComponenteCurricular.objects.filter(status='pendente').order_by('-id')
    return render(request, 'ppc/fila_aprovacao_componentes.html', {'pendentes': pendentes})


@staff_member_required(login_url='login')
def aprovar_componente(request, componente_id):
    componente = get_object_or_404(ComponenteCurricular, id=componente_id)
    if request.method == 'POST':
        componente.status = 'aprovado'
        componente.save()
        messages.success(request, f"{componente.nome} aprovado.")
    return redirect('fila_aprovacao_componentes')


@staff_member_required(login_url='login')
def rejeitar_componente(request, componente_id):
    componente = get_object_or_404(ComponenteCurricular, id=componente_id)
    if request.method == 'POST':
        with transaction.atomic():
            componente.uso_em_ppcs.all().delete()  # tira o vínculo primeiro — PROTECT exige isso
            componente.delete()
        messages.success(request, "Componente rejeitado e removido.")
    return redirect('fila_aprovacao_componentes')

@login_required
def buscar_cine_brasil_curso(request):
    termo = request.GET.get('q', '').strip()
    resultados = []
    if len(termo) >= 3:
        qs = CineBrasilCurso.objects.filter(nome_curso__icontains=termo).order_by('nome_curso')[:20]
        resultados = [
            {'id': c.id, 'nome_curso': c.nome_curso, 'area_geral_nome': c.area_geral_nome, 'rotulo_nome': c.rotulo_nome}
            for c in qs
        ]
    return JsonResponse({'resultados': resultados})


@login_required
def historico_componente(request, componente_id):
    componente = get_object_or_404(ComponenteCurricular, id=componente_id)
    registros = componente.history.all().order_by('-history_date')
    return render(request, 'ppc/historico_componente.html', {
        'componente': componente,
        'registros': registros,
    })

@login_required
def escolher_importacao_ppc(request, curso_id):
    curso = get_object_or_404(Curso, id=curso_id)
    return render(request, 'ppc/importar_ppc.html', {'curso': curso})


@login_required
def importar_ppc_modelo_novo(request, curso_id):
    """Fluxo legado: preserva a extração determinística de ``importacao.py``."""
    curso = get_object_or_404(Curso, id=curso_id)
    if request.method == 'POST':
        form = ImportarPDFForm(request.POST, request.FILES)
        if form.is_valid():
            dados = extrair_dados_pdf(request.FILES['arquivo'])
            dados_ppc = {k: v for k, v in dados.items() if k in CAMPOS_PPC_VALIDOS}
            dados_ppc.setdefault('modalidade', 'presencial')
            dados_ppc.setdefault('grau_academico', 'bacharelado')
            dados_ppc.setdefault('tipo_ppc', 'novo')
            dados_ppc.setdefault('carga_horaria_total', 0)
            dados_ppc.setdefault('numero_vagas_anuais', 0)
            dados_ppc.setdefault('duracao_minima_semestres', 0)
            dados_ppc.setdefault('duracao_media_semestres', 0)
            dados_ppc.setdefault('duracao_maxima_semestres', 0)
            dados_ppc['status'] = 'rascunho'
            ppc = PPC.objects.create(curso=curso, **dados_ppc)
            return redirect('editar_apresentacao', ppc_id=ppc.id)
    else:
        form = ImportarPDFForm()
    return render(request, 'ppc/importar_ppc_modelo_novo.html', {'form': form, 'curso': curso})

def _inteiro_nao_negativo(valor) -> int:
    """PositiveIntegerField não aceita None nem negativo; dado corrompido vira 0
    para revisão humana posterior, igual ao tratamento de campo ausente."""
    try:
        numero = int(valor)
    except (TypeError, ValueError):
        return 0
    return numero if numero >= 0 else 0


@login_required
def importar_ppc_modelo_antigo(request, curso_id):
    """Caminho 2: cria PPC em rascunho com extração por IA, com revisão humana obrigatória depois."""
    curso = get_object_or_404(Curso, id=curso_id)
    if request.method == 'POST':
        form = ImportarPPCModeloAntigoForm(request.POST, request.FILES)
        if not form.is_valid():
            return _resposta_erro_importacao(request, form, curso, "Revise os dados informados.", 400)

        chave = f"importacao_antiga_em_andamento_{curso.id}"
        if request.session.get(chave):
            return _resposta_erro_importacao(request, form, curso, "Já existe uma importação em andamento para este curso.", 409)
        request.session[chave] = True
        request.session.modified = True

        arquivo_temporario = None
        try:
            # analisar_pdf_adaptativo espera um Path em disco, não o UploadedFile do Django.
            upload = form.cleaned_data['arquivo']
            with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp:
                for pedaco in upload.chunks():
                    tmp.write(pedaco)
                arquivo_temporario = Path(tmp.name)

            dados, diagnostico = analisar_pdf_adaptativo(arquivo_temporario)
            with transaction.atomic():
                ppc = criar_ppc_rascunho_a_partir_de_dados_ia(dados, curso)
        except (ErroOpenRouter, ValueError) as erro:
            logger.exception("Falha na importação antiga do PPC, curso id=%s", curso.id)
            return _resposta_erro_importacao(request, form, curso, _mensagem_amigavel_importacao(erro), 422)
        except Exception:
            logger.exception("Erro inesperado na importação antiga do PPC, curso id=%s", curso.id)
            return _resposta_erro_importacao(request, form, curso, "Não foi possível concluir a importação agora. Tente novamente em alguns minutos.", 500)
        finally:
            if arquivo_temporario is not None:
                arquivo_temporario.unlink(missing_ok=True)
            request.session.pop(chave, None)
            request.session.modified = True

        destino = redirect('editar_apresentacao', ppc_id=ppc.id).url
        if _requisicao_ajax(request):
            return JsonResponse({"ok": True, "redirect_url": destino})
        messages.success(request, "PPC importado como rascunho. Revise todos os campos antes de finalizar.")
        return redirect(destino)
    return render(request, 'ppc/importar_ppc_modelo_antigo.html', {'form': ImportarPPCModeloAntigoForm(), 'curso': curso})

CAMPOS_PPC_VALIDOS = {  # supondo que já existe algo assim no Caminho 1 — reaproveite se já tiver
    'modalidade', 'grau_academico', 'turno_funcionamento', 'carga_horaria_total',
    'numero_vagas_anuais', 'duracao_minima_semestres', 'duracao_media_semestres',
    'duracao_maxima_semestres', 'diretor', 'vice_diretor', 'coordenador_curso', 'vice_coordenador_curso',
    'numero_resolucao', 'tipo_ppc', 'publico_alvo_ead', 'ato_integracao_uab',
    'ato_credenciamento_mec', 'polos_ead', 'apresentacao_texto', 'exposicao_motivos',
    'objetivo_geral', 'objetivo_especifico', 'principios_geral',
    'principios_pratica_profissional', 'principios_formacao_tecnica',
    'principios_formacao_etica_social', 'principios_interdisciplinaridade',
    'principios_articulacao_teoria_pratica', 'perfil_curso', 'perfil_habilidades',
    'estrutura_curricular_descricao', 'estrutura_curricular_informacoes_complementares',
    'estagio', 'tcc', 'atividades_complementares', 'politicas_integrada',
    'avaliacao_ensino_aprendizagem', 'avaliacao_projeto_curso', 'qualificacao',
    'diretrizes_curriculares_nacionais_curso', 'diretrizes_curriculares_nacionais_educacao_basica',
    'diretrizes_etnico_raciais_historia_cultura_afro_indigena', 'diretrizes_educacao_direitos_humanos',
    'protecao_direitos_pessoa_transtorno_espectro_autista', 'componente_curricular_libras',
    'politicas_educacao_ambiental', 'diretrizes_formacao_professores_educacao_basica',
    'condicoes_acesso_pessoas_deficiencia_mobilidade_reduzida', 'bibliografias_ppc',
}

FALLBACKS_PPC_OBRIGATORIOS = {
    'modalidade': 'presencial',
    'grau_academico': 'bacharelado',
    'tipo_ppc': 'novo',
    'carga_horaria_total': 0,
    'numero_vagas_anuais': 0,
    'duracao_minima_semestres': 0,
    'duracao_media_semestres': 0,
    'duracao_maxima_semestres': 0,
}


def criar_ppc_rascunho_a_partir_de_dados_ia(dados: dict, curso) -> "PPC":
    dados_ppc = {k: v for k, v in dados.get('ppc', {}).items() if k in CAMPOS_PPC_VALIDOS and v is not None}
    for campo, valor_padrao in FALLBACKS_PPC_OBRIGATORIOS.items():
        dados_ppc.setdefault(campo, valor_padrao)
    dados_ppc['status'] = 'rascunho'

    ppc = PPC.objects.create(curso=curso, **dados_ppc)

    mapa_componentes = {}
    for item in dados.get('componentes_curriculares', []):
        componente = ComponenteCurricular.objects.create(
            ppc=ppc,
            nome=item.get('nome') or '',
            tipo=item.get('tipo') or '',
            natureza=item.get('natureza') or '',
            nucleo=item.get('nucleo') or '',
            periodo=_inteiro_nao_negativo(item.get('periodo')),
            carga_horaria_teorica=_inteiro_nao_negativo(item.get('carga_horaria_teorica')),
            carga_horaria_pratica=_inteiro_nao_negativo(item.get('carga_horaria_pratica')),
            carga_horaria_pcc=_inteiro_nao_negativo(item.get('carga_horaria_pcc')),
            unidade_academica_componente=item.get('unidade_academica_componente') or '',
            ementa=item.get('ementa') or '',
        )
        mapa_componentes[item.get('nome')] = componente

        for bib in item.get('bibliografias', []):
            Bibliografia.objects.create(
                componente=componente,
                tipo=bib.get('tipo') or '',
                titulo=bib.get('titulo') or '',
                autores=bib.get('autores') or '',
                editora=bib.get('editora') or '',
                cidade=bib.get('cidade') or '',
                ano=bib.get('ano'),
            )

    for item in dados.get('componentes_curriculares', []):
        origem = mapa_componentes.get(item.get('nome'))
        if origem is None:
            continue
        for relacao in item.get('relacoes', []):
            destino = mapa_componentes.get(relacao.get('componente_relacionado_nome'))
            if destino is None or destino == origem:
                continue  # componente relacionado fora do PDF, ou auto-relação — ignora
            RelacaoComponente.objects.get_or_create(
                componente=origem,
                componente_relacionado=destino,
                tipo=relacao.get('tipo') or '',
            )

    return ppc

def _requisicao_ajax(request):
    return request.headers.get('X-Requested-With') == 'XMLHttpRequest'


def _resposta_erro_importacao(request, form, curso, mensagem, status):
    if _requisicao_ajax(request):
        return JsonResponse({"ok": False, "erro": mensagem, "campos": form.errors}, status=status)
    form.add_error(None, mensagem)
    return render(request, 'ppc/importar_ppc_modelo_antigo.html', {'form': form, 'curso': curso}, status=status)


def _rascunho_importacao_antiga(request, ppc_id):
    """Obtém somente o rascunho associado ao PPC desta sessão."""
    rascunho = request.session.get(f"rascunho_importacao_antiga_{ppc_id}", {})
    return rascunho if isinstance(rascunho, dict) else {}


def _initial_rascunho(request, ppc, classe_form, secao="ppc"):
    """Sobrepõe o banco apenas na primeira apresentação do campo em revisão."""
    dados = _rascunho_importacao_antiga(request, ppc.id).get("dados", {})
    origem = dados.get(secao, {}) if isinstance(dados, dict) else {}
    campos = getattr(classe_form.Meta, "fields", [])
    return {campo: origem[campo] for campo in campos if origem.get(campo) is not None}


def _consumir_campos_rascunho(request, ppc, classe_form, secao="ppc"):
    """Após salvar uma seção, o valor manual passa a ter prioridade sobre o rascunho."""
    chave = f"rascunho_importacao_antiga_{ppc.id}"
    rascunho = _rascunho_importacao_antiga(request, ppc.id)
    dados = rascunho.get("dados", {})
    origem = dados.get(secao, {}) if isinstance(dados, dict) else {}
    for campo in getattr(classe_form.Meta, "fields", []):
        origem.pop(campo, None)
    if origem is not None:
        dados[secao] = origem
    request.session[chave] = rascunho
    request.session.modified = True


def _mensagem_amigavel_importacao(erro):
    detalhe = str(erro).lower()
    if 'timeout' in detalhe or '504' in detalhe:
        return "A análise demorou mais do que o esperado. Nenhuma alteração foi aplicada; tente novamente."
    if 'openrouter' in detalhe or 'provider' in detalhe:
        return "O serviço de análise não respondeu corretamente. Tente novamente em alguns minutos."
    if 'json' in detalhe or 'estrutura' in detalhe:
        return "O resultado da análise não pôde ser validado. Tente novamente com o PDF original."
    if 'pdf' in detalhe or 'texto extraível' in detalhe:
        return "Não foi possível ler este PDF. Confirme se o arquivo é o documento original e tente novamente."
    return "A importação falhou. Verifique o arquivo e tente novamente."

@login_required
def lista_nde(request, curso_id):
    curso = get_object_or_404(Curso, id=curso_id)
    membros = curso.membros_nde.filter(ativo=True).order_by('-funcao', 'nome')
    rascunho = _rascunho_importacao_antiga(request, request.GET.get('ppc_id', 0))
    membros_rascunho = rascunho.get('dados', {}).get('membros_nde', [])
    return render(request, 'ppc/lista_nde.html', {'curso': curso, 'membros': membros, 'membros_rascunho': membros_rascunho})


@login_required
def criar_membro_nde(request, curso_id):
    curso = get_object_or_404(Curso, id=curso_id)
    exigir_acesso_ao_curso(request.user, curso) 
    if request.method == 'POST':
        form = MembroNDEForm(request.POST)
        if form.is_valid():
            membro = form.save(commit=False)
            membro.curso = curso
            membro.save()
            return redirect('lista_nde', curso_id=curso.id)
    else:
        form = MembroNDEForm()
    return render(request, 'ppc/criar_membro_nde.html', {'form': form, 'curso': curso})


@login_required
def editar_membro_nde(request, membro_id):
    membro = get_object_or_404(MembroNDE, id=membro_id)
    if request.method == 'POST':
        form = MembroNDEForm(request.POST, instance=membro)
        if form.is_valid():
            form.save()
            return redirect('lista_nde', curso_id=membro.curso.id)
    else:
        form = MembroNDEForm(instance=membro)
    return render(request, 'ppc/editar_membro_nde.html', {'form': form, 'curso': membro.curso})


@login_required
def excluir_membro_nde(request, membro_id):
    membro = get_object_or_404(MembroNDE, id=membro_id)
    curso_id = membro.curso.id
    if request.method == 'POST':
        membro.delete()
    return redirect('lista_nde', curso_id=curso_id)


def lista_cursos(request):
    query = request.GET.get('q', '')
    cursos = cursos_do_usuario(request.user)
    if query:
        cursos = cursos.filter(
            Q(nome__icontains=query) |
            Q(unidade_academica__icontains=query) |
            Q(area_conhecimento__icontains=query)
        ) 
    return render(request, 'ppc/lista_cursos.html', {'cursos': cursos, 'query': query})

@login_required
def excluir_apendice(request, apendice_id):
    apendice = get_object_or_404(Apendice, id=apendice_id)
    ppc_id = apendice.ppc.id
    if request.method == 'POST':
        apendice.delete()
    return redirect('editar_apendices', ppc_id=ppc_id)


@login_required
def editar_referencias(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = ReferenciasForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            return redirect('editar_referencias', ppc_id=ppc.id)
    else:
        form = ReferenciasForm(instance=ppc, initial=_initial_rascunho(request, ppc, ReferenciasForm))
    return render(request, 'ppc/editar_referencias.html', {'form': form, 'ppc': ppc})


@login_required
def editar_relacao(request, relacao_id):
    relacao = get_object_or_404(RelacaoComponente, id=relacao_id)
    componente_na_matriz = relacao.componente_na_matriz
    if request.method == 'POST':
        form = RelacaoComponenteForm(
            request.POST, instance=relacao,
            ppc=componente_na_matriz.ppc, componente_na_matriz_atual=componente_na_matriz,
        )
        if form.is_valid():
            form.save()
            return redirect('detalhe_componente_na_matriz', componente_na_matriz_id=componente_na_matriz.id)
    else:
        form = RelacaoComponenteForm(
            instance=relacao,
            ppc=componente_na_matriz.ppc, componente_na_matriz_atual=componente_na_matriz,
        )
    return render(request, 'ppc/editar_relacao.html', {
        'form': form, 'componente_na_matriz': componente_na_matriz, 'ppc': componente_na_matriz.ppc,
    })


@login_required
def excluir_relacao(request, relacao_id):
    relacao = get_object_or_404(RelacaoComponente, id=relacao_id)
    componente_na_matriz_id = relacao.componente_na_matriz.id
    if request.method == 'POST':
        relacao.delete()
    return redirect('detalhe_componente_na_matriz', componente_na_matriz_id=componente_na_matriz_id)


@login_required
def editar_bibliografia(request, bibliografia_id):
    bib = get_object_or_404(Bibliografia, id=bibliografia_id)
    componente_na_matriz = bib.componente_na_matriz
    if request.method == 'POST':
        form = BibliografiaForm(request.POST, instance=bib)
        if form.is_valid():
            form.save()
            return redirect('detalhe_componente_na_matriz', componente_na_matriz_id=componente_na_matriz.id)
    else:
        form = BibliografiaForm(instance=bib)
    return render(request, 'ppc/editar_bibliografia.html', {
        'form': form,
        'componente_na_matriz': componente_na_matriz,
        'ppc': componente_na_matriz.ppc,
    })


@login_required
def excluir_bibliografia(request, bibliografia_id):
    bib = get_object_or_404(Bibliografia, id=bibliografia_id)
    componente_na_matriz_id = bib.componente_na_matriz.id
    if request.method == 'POST':
        bib.delete()
    return redirect('detalhe_componente_na_matriz', componente_na_matriz_id=componente_na_matriz_id)


@login_required
def meus_componentes_pendentes(request):
    pendentes = ComponenteCurricular.objects.filter(
        criado_por=request.user, status='pendente'
    ).order_by('-id')
    return render(request, 'ppc/meus_componentes_pendentes.html', {'pendentes': pendentes})

@login_required
def criar_componente(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form_componente = ComponenteCurricularForm(request.POST)
        form_vinculo = ComponenteNaMatrizForm(request.POST, ppc=ppc)
        if form_componente.is_valid() and form_vinculo.is_valid():
            componente = form_componente.save(commit=False)
            componente.status = 'aprovado' if request.user.is_staff else 'pendente'
            componente.criado_por = request.user
            componente.save()
            vinculo = form_vinculo.save(commit=False)
            vinculo.ppc = ppc
            vinculo.componente = componente
            vinculo.save()
            if componente.status == 'pendente':
                messages.info(request, "Componente enviado para aprovação de um administrador — ficará oculto até ser aprovado.")
                return redirect('lista_componentes', ppc_id=ppc.id)
            return redirect('detalhe_componente_na_matriz', componente_na_matriz_id=vinculo.id)
    else:
        form_componente = ComponenteCurricularForm()
        form_vinculo = ComponenteNaMatrizForm(ppc=ppc)
    return render(request, 'ppc/criar_componente.html', {
        'form_componente': form_componente, 'form_vinculo': form_vinculo, 'ppc': ppc,
    })


@login_required
def editar_componente(request, componente_id):
    if not request.user.is_staff:
        return HttpResponseForbidden("Somente administradores podem editar componentes.")
    componente = get_object_or_404(ComponenteCurricular, id=componente_id)
    componente_na_matriz_id_retorno = request.GET.get('retorno') or request.POST.get('retorno')

    if request.method == 'POST':
        form = ComponenteCurricularForm(request.POST, instance=componente)
        if form.is_valid():
            componente = form.save(commit=False)
            componente.status = 'aprovado'  # edição de admin é sempre automática, mesmo se estava pendente
            componente.save()
            if componente_na_matriz_id_retorno:
                return redirect('detalhe_componente_na_matriz', componente_na_matriz_id=componente_na_matriz_id_retorno)
            return redirect('historico_componente', componente_id=componente.id)
    else:
        form = ComponenteCurricularForm(instance=componente)
    return render(request, 'ppc/editar_componente.html', {
        'form': form, 'componente': componente, 'retorno': componente_na_matriz_id_retorno,
    })


@login_required
def remover_componente_da_matriz(request, componente_na_matriz_id):
    componente_na_matriz = get_object_or_404(ComponenteNaMatriz, id=componente_na_matriz_id)
    ppc_id = componente_na_matriz.ppc.id
    if request.method == 'POST':
        componente_na_matriz.delete()
        return redirect('lista_componentes', ppc_id=ppc_id)
    return render(request, 'ppc/excluir_componente.html', {
        'componente_na_matriz': componente_na_matriz,
        'ppc': componente_na_matriz.ppc,
    })

@login_required
def editar_vinculo_componente(request, componente_na_matriz_id):
    componente_na_matriz = get_object_or_404(ComponenteNaMatriz, id=componente_na_matriz_id)
    if request.method == 'POST':
        form = ComponenteNaMatrizForm(request.POST, instance=componente_na_matriz)
        if form.is_valid():
            form.save()
            return redirect('detalhe_componente_na_matriz', componente_na_matriz_id=componente_na_matriz.id)
    else:
        form = ComponenteNaMatrizForm(instance=componente_na_matriz)
    return render(request, 'ppc/editar_vinculo_componente.html', {
        'form': form,
        'componente_na_matriz': componente_na_matriz,
        'ppc': componente_na_matriz.ppc,
    })


@login_required
def detalhe_componente_na_matriz(request, componente_na_matriz_id):
    componente_na_matriz = get_object_or_404(ComponenteNaMatriz, id=componente_na_matriz_id)
    componente = componente_na_matriz.componente
    if componente.status != 'aprovado' and not request.user.is_staff:
        raise Http404("Componente ainda não aprovado.")
    bibliografia_form = BibliografiaForm()
    relacao_form = RelacaoComponenteForm(ppc=componente_na_matriz.ppc, componente_na_matriz_atual=componente_na_matriz)

    if request.method == 'POST':
        if 'adicionar_bibliografia' in request.POST:
            bibliografia_form = BibliografiaForm(request.POST)
            if bibliografia_form.is_valid():
                bib = bibliografia_form.save(commit=False)
                bib.componente_na_matriz = componente_na_matriz
                bib.save()
                return redirect('detalhe_componente_na_matriz', componente_na_matriz_id=componente_na_matriz.id)
        elif 'adicionar_relacao' in request.POST:
            relacao_form = RelacaoComponenteForm(request.POST, ppc=componente_na_matriz.ppc, componente_na_matriz_atual=componente_na_matriz)
            if relacao_form.is_valid():
                relacao = relacao_form.save(commit=False)
                relacao.componente_na_matriz = componente_na_matriz
                try:
                    relacao.save()
                    return redirect('detalhe_componente_na_matriz', componente_na_matriz_id=componente_na_matriz.id)
                except ValidationError as exc:
                    relacao_form.add_error(None, exc)

    bibliografias_basicas = componente_na_matriz.bibliografias.filter(tipo='basica')
    bibliografias_complementares = componente_na_matriz.bibliografias.filter(tipo='complementar')

    return render(request, 'ppc/detalhe_componente.html', {
        'ppc': componente_na_matriz.ppc,
        'componente_na_matriz': componente_na_matriz,
        'componente': componente,
        'bibliografia_form': bibliografia_form,
        'relacao_form': relacao_form,
        'bibliografias_basicas': bibliografias_basicas,
        'bibliografias_complementares': bibliografias_complementares,
        'basicas_na_biblioteca': bibliografias_basicas.filter(referencia__isnull=False).values('referencia_id').distinct().count(),
        'complementares_na_biblioteca': bibliografias_complementares.filter(referencia__isnull=False).values('referencia_id').distinct().count(),
    })

@login_required
def buscar_componente_existente(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    termo = request.GET.get('q', '').strip()
    ja_usados_ids = set(ppc.matriz_componentes.values_list('componente_id', flat=True))
    resultados = []
    if termo:
        resultados = (ComponenteCurricular.objects.filter(nome__icontains=termo, status='aprovado').exclude(id__in=ja_usados_ids).annotate(
                        prioridade_ccu=Case(
                            When(ccu=True, then=Value(0)),
                            When(ccu=None, then=Value(1)),
                            When(ccu=False, then=Value(2)),
                            output_field=IntegerField(),
                        )
                    ).order_by('prioridade_ccu', 'nome')
                    )
    return render(request, 'ppc/buscar_componente_existente.html', {
        'ppc': ppc, 'termo': termo, 'resultados': resultados,
    })

@login_required
def adicionar_componente_existente(request, ppc_id, componente_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    componente = get_object_or_404(ComponenteCurricular, id=componente_id)
    if ComponenteNaMatriz.objects.filter(ppc=ppc, componente=componente).exists():
        messages.warning(request, "Este componente já está na matriz deste PPC.")
        return redirect('lista_componentes', ppc_id=ppc.id)

    if request.method == 'POST':
        form = ComponenteNaMatrizForm(request.POST, ppc=ppc)
        if form.is_valid():
            vinculo = form.save(commit=False)
            vinculo.ppc = ppc
            vinculo.componente = componente
            vinculo.save()
            return redirect('detalhe_componente_na_matriz', componente_na_matriz_id=vinculo.id)
    else:
        form = ComponenteNaMatrizForm(ppc=ppc)
    return render(request, 'ppc/adicionar_componente_existente.html', {
        'ppc': ppc, 'componente': componente, 'form': form,
    })

@login_required
def lista_componentes(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)

    if request.method == "POST" and "salvar_ppc_anterior" in request.POST:
        anterior_id = request.POST.get("ppc_anterior") or None
        anterior = get_object_or_404(PPC, pk=anterior_id, curso_id=ppc.curso_id) if anterior_id else None
        ppc.ppc_anterior = anterior
        try:
            ppc.save(update_fields=["ppc_anterior", "atualizado_em"])
        except ValidationError as erro:
            messages.error(request, " ".join(erro.messages))
        else:
            messages.success(request, "PPC anterior atualizado.")
            return redirect("lista_componentes", ppc_id=ppc.pk)

    if request.method == 'POST' and 'salvar_descricao' in request.POST:
        form = EstruturaCurricularForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            return redirect('lista_componentes', ppc_id=ppc.id)
    else:
        form = EstruturaCurricularForm(instance=ppc)

    resumo_manual, _ = ResumoCargaHorariaPPC.objects.get_or_create(ppc=ppc)

    if request.method == 'POST' and 'salvar_resumo_manual' in request.POST:
        form_resumo_manual = ResumoCargaHorariaPPCForm(request.POST, instance=resumo_manual)
        if form_resumo_manual.is_valid():
            form_resumo_manual.save()
            return redirect('lista_componentes', ppc_id=ppc.id)
    else:
        form_resumo_manual = ResumoCargaHorariaPPCForm(instance=resumo_manual)

    componentes_na_matriz = ppc.matriz_componentes.select_related('componente').filter(
        componente__status='aprovado'
    ).order_by('periodo', 'componente__nome')

    soma_componentes = sum(cm.componente.carga_horaria_computada_total for cm in componentes_na_matriz)
    diferenca_carga_horaria = ppc.carga_horaria_total - soma_componentes

    resumo = sincronizar_carga_horaria_total(ppc)

    return render(request, 'ppc/lista_componentes.html', {
        'ppc': ppc,
        'componentes_na_matriz': componentes_na_matriz,
        'form': form,
        'form_resumo_manual': form_resumo_manual,
        'soma_componentes': soma_componentes,
        'diferenca_carga_horaria': diferenca_carga_horaria,
        'resumo': resumo,
        'ppcs_anteriores': PPC.objects.filter(curso_id=ppc.curso_id).exclude(pk=ppc.pk).order_by('-ano_vigencia', '-pk'),
    })


@login_required
def editar_apendices(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = ApendiceForm(request.POST, request.FILES)
        if form.is_valid():
            apendice = form.save(commit=False)
            apendice.ppc = ppc
            apendice.save()
            return redirect('editar_apendices', ppc_id=ppc.id)
    else:
        form = ApendiceForm()
    apendices_rascunho = _rascunho_importacao_antiga(request, ppc.id).get('dados', {}).get('apendices', [])
    return render(request, 'ppc/apendices.html', {'form': form, 'ppc': ppc, 'apendices_rascunho': apendices_rascunho})

@login_required
def editar_dinamicas_ead(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    dinamica = DinamicaEAD.objects.filter(ppc=ppc).first() or DinamicaEAD(ppc=ppc)
    if request.method == 'POST':
        form = DinamicaEADForm(request.POST, instance=dinamica)
        if form.is_valid():
            form.save()
            return redirect('editar_dinamicas_ead', ppc_id=ppc.id)
    else:
        form = DinamicaEADForm(instance=dinamica, initial=_initial_rascunho(request, ppc, DinamicaEADForm, 'dinamica_ead'))
    return render(request, 'ppc/editar_dinamicas_ead.html', {'form': form, 'ppc': ppc})
@login_required
def editar_requisitos_legais(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = RequisitosLegaisForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            return redirect('editar_requisitos_legais', ppc_id=ppc.id)
    else:
        form = RequisitosLegaisForm(instance=ppc, initial=_initial_rascunho(request, ppc, RequisitosLegaisForm))
    return render(request, 'ppc/editar_requisitos_legais.html', {'form': form, 'ppc': ppc})

@login_required
def editar_qualificacao(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = QualificacaoForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            return redirect('editar_qualificacao', ppc_id=ppc.id)
    else:
        form = QualificacaoForm(instance=ppc, initial=_initial_rascunho(request, ppc, QualificacaoForm))
    return render(request, 'ppc/editar_qualificacao.html', {'form': form, 'ppc': ppc})

@login_required
def editar_avaliacao_projeto_curso(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = AvalicaoProjetoCursoForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            return redirect('editar_avaliacao_projeto_curso', ppc_id=ppc.id)
    else:
        form = AvalicaoProjetoCursoForm(instance=ppc, initial=_initial_rascunho(request, ppc, AvalicaoProjetoCursoForm))
    return render(request, 'ppc/editar_avaliacao_projeto_curso.html', {'form': form, 'ppc': ppc})

@login_required
def editar_avaliacao_ensino(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = AvaliacaoEnsinoForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            return redirect('editar_avaliacao_ensino', ppc_id=ppc.id)
    else:
        form = AvaliacaoEnsinoForm(instance=ppc, initial=_initial_rascunho(request, ppc, AvaliacaoEnsinoForm))
    return render(request, 'ppc/editar_avaliacao_ensino.html', {'form': form, 'ppc': ppc})

@login_required
def editar_politicas_integrada(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = PoliticasIntegradaForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            return redirect('editar_politicas_integrada', ppc_id=ppc.id)
    else:
        form = PoliticasIntegradaForm(instance=ppc, initial=_initial_rascunho(request, ppc, PoliticasIntegradaForm))
    return render(request, 'ppc/editar_politicas_integrada.html', {'form': form, 'ppc': ppc})

@login_required
def editar_atividades_complementares(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = AtividadesComplementaresForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            return redirect('editar_atividades_complementares', ppc_id=ppc.id)
    else:
        form = AtividadesComplementaresForm(instance=ppc, initial=_initial_rascunho(request, ppc, AtividadesComplementaresForm))
    return render(request, 'ppc/editar_atividades_complementares.html', {'form': form, 'ppc': ppc})

@login_required
def editar_estagio(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = EstagioForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            return redirect('editar_estagio', ppc_id=ppc.id)
    else:
        form = EstagioForm(instance=ppc, initial=_initial_rascunho(request, ppc, EstagioForm))
    return render(request, 'ppc/editar_estagio.html', {'form': form, 'ppc': ppc})

@login_required
def editar_tcc(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = TccForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            return redirect('editar_tcc', ppc_id=ppc.id)
    else:
        form = TccForm(instance=ppc, initial=_initial_rascunho(request, ppc, TccForm))
    return render(request, 'ppc/editar_tcc.html', {'form': form, 'ppc': ppc})

@login_required
def editar_expectativas(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = ExpectativasForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            return redirect('editar_expectativas', ppc_id=ppc.id)
    else:
        form = ExpectativasForm(instance=ppc, initial=_initial_rascunho(request, ppc, ExpectativasForm))
    return render(request, 'ppc/editar_expectativas.html', {'form': form, 'ppc': ppc})

@login_required
def editar_principios(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = PrincipiosForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            return redirect('editar_principios', ppc_id=ppc.id)
    else:
        form = PrincipiosForm(instance=ppc, initial=_initial_rascunho(request, ppc, PrincipiosForm))
    return render(request, 'ppc/editar_principios.html', {'form': form, 'ppc': ppc})



@login_required
def editar_apresentacao(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = ApresentacaoForm(request.POST, instance=ppc, usuario=request.user)
        if form.is_valid():
            form.save()
            _consumir_campos_rascunho(request, ppc, ApresentacaoForm)
            return redirect('editar_apresentacao', ppc_id=ppc.id)
    else:
        form = ApresentacaoForm(instance=ppc, usuario=request.user, initial=_initial_rascunho(request, ppc, ApresentacaoForm))
    return render(request, 'ppc/editar_apresentacao.html', {'form': form, 'ppc': ppc})


@login_required
def editar_exposicao_motivos(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = ExposicaoMotivosForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            _consumir_campos_rascunho(request, ppc, ExposicaoMotivosForm)
            return redirect('editar_exposicao_motivos', ppc_id=ppc.id)
    else:
        form = ExposicaoMotivosForm(instance=ppc, initial=_initial_rascunho(request, ppc, ExposicaoMotivosForm))
    return render(request, 'ppc/editar_exposicao_motivos.html', {'form': form, 'ppc': ppc})

@staff_member_required(login_url='login')
def criar_curso(request):
    if request.method == 'POST':
        form = CursoForm(request.POST)
        if form.is_valid():
            curso = form.save()
            return redirect('detalhe_curso', curso_id=curso.id)
    else:
        form = CursoForm()
    return render(request, 'ppc/criar_curso.html', {'form': form})

@staff_member_required
def editar_permissoes(request, user_id):
    usuario = get_object_or_404(User, id=user_id)
    if request.method == 'POST':
        form = EditarPermissoesForm(request.POST, instance=usuario)
        if form.is_valid():
            form.save()
            return redirect('gestao_usuarios')
    else:
        form = EditarPermissoesForm(instance=usuario)
    return render(request, 'ppc/editar_permissoes.html', {'form': form, 'usuario': usuario})




def home(request):
    total_cursos = Curso.objects.count()
    total_ppcs = PPC.objects.count()
    ppcs_recentes = PPC.objects.select_related('curso').order_by('-atualizado_em')[:5]
    return render(request, "ppc/home.html", {
        'total_cursos': total_cursos,
        'total_ppcs': total_ppcs,
        'ppcs_recentes': ppcs_recentes,
    })

def ajuda(request):
    return render(request, 'ppc/ajuda.html')

@staff_member_required
def gestao_usuarios(request):
    query = request.GET.get('q', '')
    usuarios = User.objects.all().order_by('username')
    if query:
        usuarios = usuarios.filter(
            Q(username__icontains=query) | Q(email__icontains=query)
        )
    return render(request, 'ppc/gestao_usuarios.html', {'usuarios': usuarios, 'query': query})

@staff_member_required
def criar_usuario(request):
    if request.method == 'POST':
        form = UserCreationForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('gestao_usuarios')
    else:
        form = UserCreationForm()
    return render(request, 'ppc/criar_usuario.html', {'form': form})

@staff_member_required
def alternar_acesso_usuario(request, user_id):
    usuario = get_object_or_404(User, id=user_id)
    if request.method == 'POST' and usuario != request.user:
        usuario.is_active = not usuario.is_active
        usuario.save()
    return redirect('gestao_usuarios')

@login_required
def detalhe_curso(request, curso_id):
    curso = get_object_or_404(Curso, id=curso_id)
    ppcs = curso.ppcs.all()
    return render(request, 'ppc/detalhe_curso.html', {'curso': curso, 'ppcs': ppcs})


@login_required
def criar_ppc(request, curso_id):
    curso = get_object_or_404(Curso, id=curso_id)
    if request.method == 'POST':
        form = ApresentacaoForm(request.POST, curso=curso, usuario=request.user)
        if form.is_valid():
            ppc = form.save(commit=False)
            ppc.curso = curso
            ppc.save()
            return redirect('editar_apresentacao', ppc_id=ppc.id)
    else:
        form = ApresentacaoForm(curso=curso, usuario=request.user)
    return render(request, 'ppc/criar_ppc.html', {'form': form, 'curso': curso})


@login_required
def editar_objetivos(request, ppc_id):
    ppc = get_object_or_404(PPC, id=ppc_id)
    if request.method == 'POST':
        form = ObjetivosForm(request.POST, instance=ppc)
        if form.is_valid():
            form.save()
            _consumir_campos_rascunho(request, ppc, ObjetivosForm)
            return redirect('editar_objetivos', ppc_id=ppc.id)
    else:
        form = ObjetivosForm(instance=ppc, initial=_initial_rascunho(request, ppc, ObjetivosForm))
    return render(request, 'ppc/editar_objetivos.html', {'form': form, 'ppc': ppc})


@login_required
def gerar_pdf_ppc(request, ppc_id):
    """Gera uma representação de impressão independente das telas de edição."""
    bibliografias = Bibliografia.objects.order_by('tipo', 'autores', 'titulo')
    relacoes = RelacaoComponente.objects.select_related(
        'componente_relacionado_na_matriz__componente'
    ).order_by('tipo', 'componente_relacionado_na_matriz__componente__nome')

    componentes_na_matriz = ComponenteNaMatriz.objects.filter(
        componente__status='aprovado'
    ).select_related('componente').order_by('periodo', 'componente__nome').prefetch_related(
        Prefetch('bibliografias', queryset=bibliografias),
        Prefetch('relacoes', queryset=relacoes),
    )

    ppc = get_object_or_404(
        PPC.objects.select_related('curso').prefetch_related(
            Prefetch('matriz_componentes', queryset=componentes_na_matriz),
            Prefetch('apendices', queryset=Apendice.objects.order_by('tipo', 'titulo')),
        ),
        id=ppc_id,
    )

    # A relação é opcional: PPCs presenciais não precisam ter DinamicaEAD.
    dinamica_ead = DinamicaEAD.objects.filter(ppc=ppc).first()
    html = render_to_string('ppc/pdf/ppc_documento.html', {
        'ppc': ppc,
        'componentes_na_matriz': ppc.matriz_componentes.all(),
        'dinamica_ead': dinamica_ead,
    }, request=request)

    pdf = HTML(string=html, base_url=request.build_absolute_uri('/')).write_pdf()
    registrar_evento("pdf_gerado", usuario=request.user, ppc=ppc)
    nome = slugify(ppc.curso.nome) or f'ppc-{ppc.pk}'
    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="ppc-{nome}.pdf"'
    return response
